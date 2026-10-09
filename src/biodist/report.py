"""The study report: the sections the user ticks, in their order, as paragraphs and tables.

No Qt here. `to_markdown` gives one readable file (the window also renders it to .odt,
.pdf and .html through Qt); `to_xlsx` gives one workbook with a sheet per section, the
numbers left as numbers.
"""

from __future__ import annotations

import datetime as _dt
import statistics
from pathlib import Path

from . import __version__
from . import study as _study
from .study import (
    ARRIVE_NEEDED, BASES, EVENT_NEEDED, KIND_LABEL, PRESET_LABEL, Result, Study, arrive_gaps,
    at_imaging, at_injection, event_text, half_life_s, imaging_label, life_dates, recovery_pct,
    time_text,
)

SECTIONS = [
    ("general", "General"),
    ("provenance", "Sources: data files and other measures"),
    ("animals", "Animals"),
    ("arrive", "ARRIVE check"),
    ("values", "Tissue values"),             # one per unit: the one kind with options, copied
    ("checks", "Checks"),
    ("raw", "Raw counts"),
    ("agreement", "Repeated measures: how well they agree"),
]
# the tissue-value units; SUV here is body-weight SUV, (Bq/g) / (injected Bq / body g): g/g
VALUE_UNITS = [("pid_g", "%IA/g"), ("pid", "%IA"), ("suv", "SUV (g/g)"),
               ("mbq", "MBq"), ("kbq", "kBq"), ("bq_g", "Bq/g"), ("kbq_g", "kBq/g"),
               ("mass", "mass (g)"), ("mass_mg", "mass (mg)")]
LAYOUTS = [("tissues", "tissues down, animals across"), ("animals", "animals down"),
           ("list", "comma list, a line per tissue")]
# each section's options: key -> [(option, kind, label, choices, default)]; "multi" = ticks
OPTIONS = {
    "general": [("parts", "multi", "show", [
        ("study", "the study: animals, tracer, counter, window, efficiency"),
        ("picks", "the values chosen by hand (Results), if any"),
        ("processing", "the rules the values were made by"),
        ("software", "software, report date")], ["study", "picks", "software"])],
    "provenance": [("files", "choice", "data files as", [("table", "a table"),
                                                         ("list", "a list of names")], "table"),
                   ("which", "choice", "which files", [
                       ("used", "those a value came from"),
                       ("all", "all, with what each gave")], "used"),
                   ("started", "bool", "with each file's start time", None, False),
                   ("sha", "bool", "with each file's fingerprint (SHA-256)", None, False),
                   ("typed", "choice", "other measures as",
                    [("statement", "a sentence"), ("inline", "a sentence, the values in a line"),
                     ("table", "a table")], "inline")],
    "animals": [("parts", "multi", "show", [("animals", "the animals"),
                                            ("before", "procedures before the study day"),
                                            ("day", "the study day: tracer, procedures")],
                 ["animals", "before", "day"])],
    "arrive": [("detail", "choice", "list", [("missing", "what is missing"),
                                             ("all", "every item, per animal")], "missing")],
    "checks": [("flagged", "bool", "the values flagged", None, True),
               ("log", "bool", "what the checks found, in order", None, True)],
    "agreement": [("side", "bool", "every value measured more than once, its versions side "
                                   "by side", None, False)],
}
DEFAULT = [{"key": k, "on": k in ("general", "provenance", "animals", "arrive", "checks")}
           for k, _ in SECTIONS if k != "values"]
DEFAULT[4:4] = [{"key": "values", "on": u in ("pid_g", "mass"), "unit": u, "layout": "tissues",
                 "digits": None} for u in ("pid_g", "mass", "pid", "suv", "mbq")]
NEW_ON = {"checks"}       # a section added since: on in a list saved before it existed


def clean(specs) -> list[dict]:
    """The report's sections as saved in Options, checked: unknown ones dropped, a fixed one
    missing added back (off, but for a new one), options defaulted. A section removed is
    kept, "gone" (hidden, off), so it does not come back; a copy is one more of its key."""
    out = [sp for x in (specs if isinstance(specs, list) else []) if (sp := _one(x))]
    fresh, have = not out, {y["key"] for y in out}
    return out + [_one(dict(d, on=d["on"] if fresh else d["key"] in NEW_ON))
                  for d in DEFAULT if d["key"] not in have]


def _one(x) -> dict | None:
    """One section's spec with its options checked; None if it is no section."""
    if not isinstance(x, dict) or x.get("key") not in dict(SECTIONS):
        return None
    gone = {"gone": True} if x.get("gone") is True else {}
    on = bool(x.get("on")) and not gone
    if x["key"] == "values":
        d = x.get("digits")
        return {"key": "values", "on": on,
                "unit": x.get("unit") if x.get("unit") in dict(VALUE_UNITS) else "pid_g",
                "layout": x.get("layout") if x.get("layout") in dict(LAYOUTS) else "tissues",
                "digits": d if isinstance(d, int) and 0 <= d <= 8 else None, **gone}
    sp = {"key": x["key"], "on": on, **gone}
    for name, kind, _, choices, default in OPTIONS.get(x["key"], []):
        v = x.get(name, default)
        ok = (isinstance(v, bool) if kind == "bool" else
              v in dict(choices) if kind == "choice" else
              isinstance(v, list) and all(c in dict(choices) for c in v))
        sp[name] = v if ok else default
    return sp


def title(spec) -> str:
    if spec["key"] != "values":
        return dict(SECTIONS)[spec["key"]]
    u = spec.get("unit", "pid_g")
    return {"mass": "Tissue masses (g)", "mass_mg": "Tissue masses (mg)",
            "mbq": "Activity in the tissues (MBq)", "kbq": "Activity in the tissues (kBq)",
            "kbq_g": "Activity concentration (kBq/g)",
            "bq_g": "Activity concentration (Bq/g)"}.get(u, f"Tissue uptake, {dict(VALUE_UNITS)[u]}")


def _hhmm(t) -> str:
    return f"{t:%Y-%m-%d %H:%M}" if t else ""


def _grid(study: Study, res: Result, fn, summary=True):
    """Tissues down, animals across, for the collected tissues."""
    ids = [a.id for a in study.output_animals()]
    rows = [[study.tissue_label(t.name)] + [fn(a, t.name) for a in ids]
            for t in study.output_tissues()]
    if summary:
        rows.append(["injected activity (MBq, at injection)"]
                    + [_mbq(at_injection(study, res, a, res.injected_bq.get(a))) for a in ids])
    return ["tissue"] + [study.column_label(a) for a in study.output_animals()], rows


def _value(study: Study, res: Result, a: str, t: str, unit: str):
    c = res.cell(a, t)
    if unit in ("mass_mg", "mbq"):
        v = c.mass_g if unit == "mass_mg" else res.at_ref(a, c.bq)
        return None if v is None else v * 1000 if unit == "mass_mg" else v / 1e6
    return res.value(study, a, t, unit)


def _values(study: Study, res: Result, unit: str, layout: str, nd: int) -> list:
    """One unit's tissue values, laid out as asked, `nd` decimals."""
    hd, rows = _grid(study, res, lambda a, t: _value(study, res, a, t, unit),
                     not unit.startswith("mass"))
    if unit == "pid":
        rows.append(["sum of tissues (%IA)"] + [recovery_pct(study, res, a.id)
                                                for a in study.output_animals()])
    rows = [[r[0]] + [None if v is None else round(v, nd) for v in r[1:]] for r in rows]
    if layout == "animals":
        return [("table", ["animal"] + [r[0] for r in rows],
                 [[h] + [r[j] for r in rows] for j, h in enumerate(hd[1:], start=1)], nd)]
    if layout == "list":
        return [("p", "animals: " + ", ".join(hd[1:]))] + [
            ("p", f"{r[0]}: " + ", ".join("–" if v is None else f"{v:.{nd}f}" for v in r[1:]))
            for r in rows]
    return [("table", hd, rows, nd)]


def _all(names, every) -> str:
    return "(all)" if not names or names == every else ", ".join(names)


def _mbq(bq):
    return None if bq is None else bq / 1e6


def _spread(vals) -> float | None:
    vals = [v for v in vals if v is not None and v > 0]
    return (max(vals) - min(vals)) / min(vals) * 100.0 if len(vals) > 1 else None


def _agreement(study: Study, res: Result, side: bool) -> list:
    """Where a value exists more than once — a vial counted twice, read in several windows,
    a tissue weighed two ways — how far apart the versions are."""
    warn = study.count_tol_pct                        # the study's agreement
    per = {"repeat countings": {}, "energy windows": {}, "weighings": {}}
    for key, c in res.cells.items():
        per["repeat countings"][key] = c.spread_pct(study)
        per["energy windows"][key] = _spread(b for src, _, b in c.windows if src == c.bq_src)
        per["weighings"][key] = _spread(g for _, g in c.mass_alts)
    summary, over = [], {}
    for what, d in per.items():
        v = [x for x in d.values() if x is not None]
        summary.append([what, len(v), statistics.median(v) if v else None, max(v, default=None),
                        sum(1 for x in v if x > warn)])
        for key, x in d.items():
            if x is not None and x > warn:
                over.setdefault(key, {})[what] = x
    out = [("p", "Some values were measured more than once: a vial counted in several rounds, "
                 "a counting read in several energy windows (once corrected for dead time they "
                 "should give the same Bq), a tissue weighed more than one way. The spread is "
                 "how far apart a value's versions are: (largest − smallest) / smallest, in %. "
                 f"Only one version is used (the side panel of the Results says which); over "
                 f"{warn:g} % the cell is listed."),
           ("table", ["compared", "cells", "median %", "max %", f"over {warn:g}%"], summary)]
    if over:
        out.append(("table", ["animal", "tissue", *per],
                    [[a, t, *(over[(a, t)].get(w) for w in per)] for a, t in sorted(over)]))
    if side:
        rows = []
        for (a, t), c in sorted(res.cells.items()):
            if len(c.alts) > 1:
                rows.append([a, t, "activity (Bq)", "; ".join(
                    f"{Path(s).stem[:24]}: {b:,.0f}" + (" ◀" if s == c.bq_src else "")
                    for s, b, *_ in c.alts)])
            if len(c.mass_alts) > 1:
                rows.append([a, t, "mass (g)", "; ".join(
                    f"{Path(s).stem[:24]}: {g:.4f}" + (" ◀" if s == c.mass_src else "")
                    for s, g in c.mass_alts)])
        out += [("p", "Every value with more than one version, side by side (◀ the one used):"),
                ("table", ["animal", "tissue", "value", "versions"], rows)]
    return out


def _closed(study: Study, m) -> str:
    t = study.tissue(m.tissue)
    gone = [w for w, v in (("mass", m.mass_g), ("activity", m.mbq)) if v is not None
            and t and w in t.closed]
    return f"{' and '.join(gone)} row closed: not used" if gone else ""


def _effs(study: Study, runs) -> str:
    """Every window of the counting files: 'Hidex AMG 2240360 — 112-168 keV: 0.781 (the
    file's), 15-2047 keV: 0.871 (the file's), used; Wizard2 — Tc-99m: 0.75 (typed), used'."""
    per: dict[str, dict[str, str]] = {}
    for s in study.sources:
        run = runs.get(s.path)
        if not run or s.kind not in ("count", "weigh_count"):
            continue
        used = study.window_for(run)
        for w, short in zip(run.windows, run.short_windows()):
            e, own = study.eff_for(run, w)
            per.setdefault(run.counter, {}).setdefault(short, (
                short + (" keV" if short[:1].isdigit() else "") + ": "
                + (f"{e:g} ({'the file' if own else 'typed'})" if e else "none — CPM only")
                + (", used" if w == used else "")))
    return "; ".join(f"{c} — " + ", ".join(ws.values()) for c, ws in per.items()) or "—"


def _bounds(lo, hi, unit) -> str:
    """'0.5–5 mg', '≤ 5 mg', '≥ 0.5 mg' (0: no bound)."""
    return f"{lo:g}–{hi:g} {unit}" if lo and hi else f"≥ {lo:g} {unit}" if lo else \
        f"≤ {hi:g} {unit}"


def _choice(study: Study, res: Result, runs) -> list:
    w = study.window or ("auto: widest if one isotope, else the photopeak"
                         if study.window_rule == "wide" else "auto: the photopeak")
    basis, top = BASES.get(study.min_basis, ""), BASES.get(study.max_basis, "")
    target = (f"≥ {study.min_counts:g} {basis}" + (f", ≤ {study.cpm_max:g} {top}"
                                                   if study.cpm_max else "")
              + f", dead time ≤ {study.dt_max:g}")
    valid = (f"≥ {study.valid_counts:g} {basis}" + (f", ≤ {study.valid_max:g} {top}"
                                                    if study.valid_max else "")
             + f", dead time ≤ {study.valid_dt:g}")
    both = "their mean" if study.combine != "weighted" else "weighted by their counts"
    pick = {"first": "the first counting in range", "last": "the last counting in range",
            "all": f"every counting in range, {both}",
            "auto": "the counting in range with the most counts"}.get(
        study.pick_count, f"the counting round of {study.pick_count[6:]}") + \
        f" (target range: {target}; none there: the valid ones — {valid}; none valid: the " \
        f"nearest, flagged)" + (
            f"; a counting out of the consensus of the valid ones (more than half agreeing "
            f"within {study.count_tol_pct:g} % or {study.count_tol_sigma:g} σ) left out first"
            if study.count_agree else "")
    hl = sorted({a.isotope for a in study.animals if a.isotope})
    return [["counting window", w],
            ["windows used", ", ".join(sorted({x for s in study.sources
                                               if s.kind in ("count", "weigh_count")
                                               and s.path in runs
                                               and (x := study.window_for(runs[s.path]))}))],
            ["activity typed by hand (dose calibrator)", "wins over the counter"
             if study.pick_bq != "files" else "not used: the counter only"],
            ["vial counted more than once", pick],
            ["several countings chosen for a cell", both],
            ["expected per tissue (flags)", "; ".join(
                f"{r[0]} " + ", ".join(_bounds(lo, hi, u) for lo, hi, u in (
                    (r[1], r[2], "mg"), (r[3], r[4], "%IA/g")) if lo or hi)
                for r in study.ranges) or "none"],
            ["tube weighed more than once", {"first": "the first (day-of) weighing",
                                             "last": "the last weighing"}.get(
                study.mass_rule, f"the {study.mass_rule} of the weighings") + (
                f"; a weighing out of the consensus (more than half agreeing within "
                f"{study.mass_tol_mg:g} mg or {study.mass_tol_pct:g} %) left out first"
                if study.mass_agree else "")],
            ["mass", ("typed by hand wins over the tubes" if study.pick_mass != "files"
                      else "typed by hand not used: the tubes only")
             + ({"scale": "; weighings scaled by their control tubes",
                 "offset": "; weighings corrected by their control tubes' change"}.get(
                     {True: "scale"}.get(study.drift_fix, study.drift_fix), ""))],
            ["injection site (tail)", "subtracted from the injected activity"
             if study.subtract_tail else "not subtracted"],
            ["half-life", ", ".join(f"{i} {h / 3600:g} h" for i in hl if (h := half_life_s(i)))],
            ["activities (Bq, MBq) at", "each animal's injection time" if res.refs
             else _hhmm(res.ref)]]


DAY_KINDS = ("pre/co-injection", "pre-injection", "co-injection", "imaging",
             "anaesthesia")                                          # on the study day
_DAY_EXTRA = ("injection route", "euthanasia", "euthanasia time", "injection volume")


def _table(heads, rows) -> tuple:
    """A table without the columns nobody filled."""
    keep = [j for j in range(len(heads)) if j == 0 or any(r[j] not in (None, "") for r in rows)]
    return ("table", [heads[j] for j in keep], [[r[j] for j in keep] for r in rows])


def _animals(study: Study, res: Result, parts) -> list:
    b, day = [], study.day
    if "animals" in parts:
        extra = [k for k in dict.fromkeys(k for a in study.animals for k in a.extra)
                 if k.lower() not in ("dob", "date of birth", "age", *_DAY_EXTRA)]
        rows = []
        for a in study.animals:
            lf = life_dates(a.extra, day)
            rows.append([a.id, ", ".join(a.aliases), f"{lf['DOB']:%Y-%m-%d}" if lf["DOB"] else "",
                         None if lf["age"] is None else round(lf["age"], 1), a.weight_g]
                        + [a.extra.get(k, "") for k in extra] + [a.note])
        b += [("h", "The animals"), _table(["ID", "aliases", "born", "age (wk)", "weight (g)"]
                                           + extra + ["note"], rows)]
    ev = [[a.label, event_text(e)] for a in study.animals for e in a.events
          if e.get("kind") not in DAY_KINDS]
    if "before" in parts and ev:
        b += [("h", "Procedures before the study day"), ("table", ["animal", "procedure"], ev)]
    if "day" in parts:
        img = imaging_label(study, res)
        rows = [[a.id, a.isotope, a.molecule, a.extra.get("injection route", ""),
                 a.extra.get("injection volume", ""), a.full_mbq, a.full_time, a.empty_mbq,
                 a.empty_time, "; ".join(f"{lo.get('mbq')} at {lo.get('time')}" + (
                     f" ({lo['what']})" if lo.get("what") else "") for lo in a.losses),
                 a.inj_time, a.tail_mbq, a.tail_time,
                 _mbq(at_injection(study, res, a.id, res.injected_bq.get(a.id))),
                 "; ".join(f"{m} {t:%H:%M}: {bq / 1e6:.3g}"
                           for m, t, bq in at_imaging(study, res, a.id)),
                 a.extra.get("euthanasia", ""),
                 time_text(a, a.extra.get("euthanasia time", ""), study.day),
                 "; ".join(a.bio_notes)]
                for a in study.animals]
        b += [("h", "The study day: tracer, injection, procedures"),
              _table(["ID", "isotope", "molecule", "route", "volume (µL)", "syringe full (MBq)",
                      "at", "syringe empty (MBq)", "at", "other losses (MBq)", "injection",
                      "tail (MBq)", "at", "injected activity (MBq, at injection)",
                      f"at {img} start (MBq)", "euthanasia", "euthanasia time", "notes"], rows)]
        ev = [[a.label, event_text(e)] for a in study.animals for e in a.events
              if e.get("kind") in DAY_KINDS]
        if ev:
            b.append(("table", ["animal", "procedure"], ev))
    return b


def _arrive(study: Study, detail: str) -> list:
    b = [("p", "What ARRIVE 2.0 asks to report per animal — item 8, the animals; item 9, each "
               "procedure done (an imaging session: its modality, start and anaesthesia…) — read "
               "off the animal cards and their procedures.")]
    if detail == "all":
        kinds = sorted({e.get("kind") for a in study.animals for e in a.events
                        if e.get("kind") in EVENT_NEEDED})
        items = [i for i, _ in ARRIVE_NEEDED] + [f"{k}: {f}" for k in kinds
                                                  for f in EVENT_NEEDED[k]]
        rows = []
        for a in study.animals:
            gaps = set(arrive_gaps(a))
            done = {f"{e.get('kind')}: {f}" for e in a.events for f in EVENT_NEEDED.get(
                e.get("kind", ""), [])}
            rows.append([a.label] + ["–" if i in gaps else "✓" if ":" not in i or i in done
                                     else "" for i in items])
        return b + [("table", ["animal", *items], rows)]
    gaps = [[a.label, ", ".join(g)] for a in study.animals if (g := arrive_gaps(a))]
    return b + ([("p", "Missing:"), ("table", ["animal", "missing"], gaps)] if gaps else
                [("p", "Nothing missing: every item is filled for every animal.")])


def _used(res: Result) -> dict[str, dict[str, int]]:
    """file -> how many values it gave: activities, masses, tares."""
    out: dict[str, dict[str, int]] = {}
    for c in res.cells.values():
        for what, files in (("activity", c.bq_used), ("mass", c.mass_used),
                            ("tare", c.mass_empty.split(" + ") if c.mass_empty else [])):
            for f in files:
                d = out.setdefault(f.rsplit("@", 1)[0], {})
                d[what] = d.get(what, 0) + 1
    return out


def _provenance(study: Study, res: Result, runs, spec) -> list:
    b = [("h", "Data files"), ("p", f"Counting efficiency (counts per decay; Bq = CPM / 60 / "
                                    f"efficiency): {_effs(study, runs)}.")]
    used = _used(res)
    srcs = [s for s in study.sources if spec["which"] == "all" or Path(s.path).name in used]
    if len(srcs) < len(study.sources):
        b.append(("p", f"The {len(srcs)} files the values came from ({len(study.sources)} "
                       "in the study)."))
    if spec["files"] == "list":
        b += [("p", f"- {Path(s.path).name}") for s in srcs]
    else:
        heads = ["file", "kind", "started", "vials", "animals", "tissues", "order", "batch",
                 "gave", "SHA-256"]
        rows = [[Path(s.path).name, KIND_LABEL.get(s.kind, s.kind),
                 _hhmm(runs[s.path].started) if s.path in runs else "not read",
                 len(runs[s.path].slots) if s.path in runs else None,
                 _all(s.animals, [a.id for a in study.animals]),
                 _all(s.tissues, [t.name for t in study.tissues]), PRESET_LABEL[s.preset],
                 s.batch or "main",
                 ", ".join(f"{n} {w}" for w, n in used.get(Path(s.path).name, {}).items())
                 or "nothing used",
                 s.stamp.get("sha256", "")[:16]] for s in srcs]
        drop = {2: not spec["started"], 7: len(study.batch_names()) < 2,
                8: spec["which"] != "all", 9: not spec["sha"]}
        keep = [j for j in range(len(heads)) if not drop.get(j)]
        b.append(("table", [heads[j] for j in keep], [[r[j] for j in keep] for r in rows]))
    typed = [m for m in study.manual if m.mass_g is not None or m.mbq is not None or m.note]
    if not typed:
        return b
    b.append(("h", "Other measures"))
    act = sorted({m.tissue for m in typed if m.mbq is not None})
    mass = sorted({m.tissue for m in typed if m.mass_g is not None})
    say = "Some values were measured apart from the counter and the tube weighings: " + "; ".join(
        x for x in (f"activities read on the dose calibrator ({', '.join(act)})" if act else "",
                    f"masses weighed apart ({', '.join(mass)})" if mass else "") if x) \
        + ". They take the place of the files' values."
    b.append(("p", say))
    if spec["typed"] == "inline":
        for tname in dict.fromkeys(m.tissue for m in typed):
            ms = [m for m in typed if m.tissue == tname]
            parts = [f"{m.animal}: " + ", ".join(x for x in (
                f"{m.mass_g:g} g" if m.mass_g is not None else "",
                f"{m.mbq:g} MBq at {m.time}" if m.mbq is not None else "", m.note) if x)
                + (f" ({c})" if (c := _closed(study, m)) else "") for m in ms]
            b.append(("p", f"{tname} — " + "; ".join(parts)))
    elif spec["typed"] == "table":
        b.append(("table", ["animal", "tissue", "mass (g)", "activity (MBq)", "read at", "note"],
                  [[m.animal, m.tissue, m.mass_g, m.mbq, m.time,
                    "; ".join(x for x in [m.note, _closed(study, m)] if x)] for m in typed]))
    return b


def build(study: Study, res: Result, runs: dict, log: list[str], specs, digits=None) -> list:
    """[(title, blocks)] for the ticked sections, in their order (see `clean`). A block is
    ("p", text), ("h", subheading) or ("table", header, rows[, decimals])."""
    digits = digits or _study.DIGITS
    out = []
    for spec in clean(specs):
        key = spec["key"]
        if not spec["on"]:
            continue
        b = []
        if key == "general":
            iso = sorted({a.isotope for a in study.animals if a.isotope})
            mol = sorted({a.molecule for a in study.animals if a.molecule})
            rows = []
            if "study" in spec["parts"]:
                rows += [["study", study.name], ["date", study.date],
                         ["animals", len(study.animals)],
                         ["tissues", sum(1 for t in study.tissues if t.role == "tissue")],
                         ["isotope", ", ".join(iso)], ["molecule", ", ".join(mol)],
                         ["counting efficiency (counts per decay)", _effs(study, runs)]]
            if "picks" in spec["parts"] and study.chosen:
                rows.append(["chosen by hand (Results)", "; ".join(
                    f"{a}/{tn} " + ("activity" if what == "count" else what) + ": "
                    + (res.source_label(a, tn, "activity" if what == "count" else "mass")
                       if what != "empty" else "tare " + ", ".join(_study._stem(x)[-15:]
                                                                   for x in src))
                    for a, tn, what, *src in study.chosen)])
            if "processing" in spec["parts"]:
                rows += _choice(study, res, runs)
            if "software" in spec["parts"]:
                rows += [["report written", _hhmm(_dt.datetime.now())],
                         ["software", f"BioDist {__version__}"]]
            b.append(("table", ["", "value"], rows))
        elif key == "provenance":
            b += _provenance(study, res, runs, spec)
        elif key == "checks":
            flagged = [[a, t, "; ".join(c.flags)] for (a, t), c in res.cells.items() if c.flags]
            if spec["flagged"] and flagged:
                b += [("p", "Values flagged:"), ("table", ["animal", "tissue", "flags"],
                                                  sorted(flagged))]
            if spec["log"] and log:
                b.append(("p", "What the checks found, in order (✓ resolved):"))
                b += [("p", f"- {x}") for x in log]
            if not b:
                b.append(("p", "Nothing flagged."))
        elif key == "animals":
            b += _animals(study, res, spec["parts"])
        elif key == "arrive":
            b += _arrive(study, spec["detail"])
        elif key == "agreement":
            b += _agreement(study, res, spec["side"])
        elif key == "values":
            nd = spec["digits"] if spec["digits"] is not None else digits.get(spec["unit"], 2)
            b += _values(study, res, spec["unit"], spec["layout"], nd)
        elif key == "raw":
            rows = []
            for s in study.sources:
                run = runs.get(s.path)
                w = study.window_for(run) if run else None
                if not run or s.kind not in ("count", "weigh_count") or not w:
                    continue
                m = s.mapping(run)
                rows += [[run.name, sl.key, *m.get(sl.key, ("", "")), _hhmm(sl.time),
                          sl.counts.get(w), sl.cpm.get(w), sl.bq.get(w), sl.dead_time,
                          _hhmm(run.normalized_to)] for sl in run.slots]
            b.append(("p", "As exported by the counter: Bq is dead-time corrected and "
                           "normalised to the file's own instant, before any correction here."))
            b.append(("table", ["file", "rack:vial", "animal", "tissue", "counted at", "counts",
                                "CPM", "Bq", "dead time", "normalised to"], rows))
        out.append((title(spec), b))
    return out


def _txt(v, nd=None) -> str:
    if v is None:
        return ""
    if isinstance(v, float):
        if nd is not None:
            return f"{v:.{nd}f}"
        return f"{v:,.0f}" if abs(v) >= 1000 else f"{v:.4g}"
    return str(v).replace("|", "/").replace("\n", " ")


def to_markdown(sections, title: str) -> str:
    md = [f"# {title}", ""]
    for name, blocks in sections:
        md += [f"## {name}", ""]
        for b in blocks:
            if b[0] in ("p", "h"):
                md += [("### " if b[0] == "h" else "") + b[1], ""]
                continue
            hd, rows, nd = b[1], b[2], (b[3] if len(b) > 3 else None)
            md += ["| " + " | ".join(_txt(h) or " " for h in hd) + " |",
                   "|" + "---|" * len(hd)]
            md += ["| " + " | ".join(_txt(v, nd) for v in r) + " |" for r in rows]
            md.append("")
    return "\n".join(md)


def to_xlsx(sections, path) -> None:
    import openpyxl
    from openpyxl.styles import Font

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, blocks in sections:
        ws = wb.create_sheet(name.replace("/", " per ")[:31].translate(
            str.maketrans("", "", "\\?*[]:")))
        for b in blocks:
            if b[0] in ("p", "h"):
                ws.append([b[1]])
                if b[0] == "h":
                    ws[ws.max_row][0].font = Font(bold=True)
            else:
                ws.append(b[1])
                for c in ws[ws.max_row]:
                    c.font = Font(bold=True)
                for r in b[2]:
                    ws.append(list(r))
                    if len(b) > 3:
                        for c in ws[ws.max_row][1:]:
                            c.number_format = "0" + ("." + "0" * b[3] if b[3] else "")
            ws.append([])
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(40, max(
                10, *(len(_txt(c.value)) + 2 for c in col if c.value is not None
                      and len(str(c.value)) < 60)))
    wb.save(path)
