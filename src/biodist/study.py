"""The BioDist data model and the %IA/g computation.

A study is a grid: rows are tissues, columns are animals.  Every cell needs a
**mass** and an **activity**, and each of those comes from a *source* — a Hidex run
or a hand-typed entry.  Sources are the only thing that differs between the
workflows people actually use:

    tubes tared before and after   an "empty tubes" and a "filled tubes" run on the same cells
    weigh-and-count in one pass    a weigh-and-count run (+ a Tare run if it did not tare)
    "I weigh my own tubes"         manual masses, plus count-only runs
    too hot for the counter        a manual dose-calibrator activity

and in how a run's vials line up with the grid — animal-major (all of animal 1's
tissues, then animal 2's) or tissue-major (all the tumours, then all the muscles).
That is the whole abstraction; nothing else is special-cased.

%IA/g is independent of the reference time, because decay cancels between the
tissue and the injected activity as long as both are corrected to the same instant.
The reference only shows up in the Bq columns.

Run: python -m biodist.study --self-check
"""

from __future__ import annotations

import csv
import datetime as _dt
import hashlib
import itertools
import json
import math
import re
import statistics
from dataclasses import asdict, dataclass, field
from pathlib import Path

from . import hidex

# Isotope -> half-life in hours. Display name first; lookup is punctuation-insensitive.
# Options edits the list: an isotope added there gets its half-life.
ISOTOPES = [
    ("99mTc", 6.00718), ("68Ga", 1.12782), ("18F", 1.8288), ("111In", 67.313),
    ("177Lu", 159.528), ("123I", 13.2234), ("125I", 1425.6), ("64Cu", 12.7004),
    ("201Tl", 73.01), ("90Y", 64.053), ("161Tb", 166.9),
]
HALF_LIFE_S = {k: v * 3600.0 for k, v in ISOTOPES}

ANIMAL_MAJOR, TISSUE_MAJOR = "animal_major", "tissue_major"
PRESET_LABEL = {ANIMAL_MAJOR: "animal by animal", TISSUE_MAJOR: "tissue by tissue"}

TISSUE_ROLES = ["tissue", "tail", "blank", "standard"]
ROLE_HELP = {
    "tissue": "a collected tissue — appears in the results",
    "tail": "injection site — its activity is subtracted from the injected activity",
    "blank": "an empty control tube — excluded from the results, checked against background",
    "standard": "an aliquot of the injectate — cross-checks the syringe-derived dose",
}

# What a source contributes. A Hidex tare run is either the empty or the filled tubes.
KINDS = ["empty", "filled", "count", "weigh_count", "ignored"]
KIND_LABEL = {"empty": "tare", "filled": "weight", "count": "count",
              "weigh_count": "count + weight", "ignored": "(ignored)"}


def kind_of(run: hidex.Run) -> str:
    """A source's first kind: a weighing is taken as empty tubes until matched."""
    return "empty" if run.kind == "tare" else run.kind


def guess_role(name: str) -> str:
    n = name.strip().strip("()").strip().lower()             # "(empty)" is an empty vial
    if n.startswith(("tail", "inj")):
        return "tail"
    if n in ("ctrl", "control", "blank", "empty", "control tube"):
        return "blank"
    if n in ("std", "standard", "injectate"):
        return "standard"
    return "tissue"

def unique_names(names, taken=()) -> list[str]:
    """Every tissue is found by its name: a name seen before gets a number, '(empty) 2'."""
    out, seen = [], set(taken)
    for n in names:
        m, k = n, 2
        while m in seen:
            m, k = f"{n} {k}", k + 1
        seen.add(m)
        out.append(m)
    return out


# Fields beyond the minimum, offered by the animal card's "+ field", and the order the card
# lists them in (related ones side by side). Free text: the list is a convenience, not a
# schema, so an unlisted ARRIVE item is just typed in and goes after these.
CARD_FIELDS = ["species", "strain", "sex", "DOB"]            # always on the card
ARRIVE_FIELDS = [
    "group", "supplier", "arrival", "age at arrival", "protocol", "housing",
    "health status", "injection route", "euthanasia", "euthanasia time",
    "exclusion reason", "comment",
]
HOUSING = ("5 per cage, 21-23 °C, 12:12 h light/dark, food and water ad libitum, enrichment "
           "(nesting material, shelter)")
FIELD_HINT = {"arrival": "the day the animals came in", "age at arrival": "6 wk, 42 d",
              "protocol": "IACUC / ethics protocol #",
              "euthanasia time": "HH:MM, or 2 h p.i.",
              "injection route": "i.v. tail vein…", "supplier": "Janvier, in-house…",
              "housing": HOUSING,
              "health status": "SPF, SOPF, conventional…",
              "euthanasia": "method: CO2, cervical dislocation under anaesthesia…",
              "age": "12 wk, 2.5 mo, 10-12 wk — instead of a date of birth",
              "injection volume": "µL"}
# what the animal window's lists offer (anything can still be typed) — Options › Animal fields
ANIMAL_LISTS = {"euthanasia": ["CO2", "cervical dislocation under anaesthesia",
                               "exsanguination under anaesthesia",
                               "decapitation under anaesthesia", "pentobarbital overdose"],
                "injection route": ["i.v. tail vein", "i.v. retro-orbital", "i.p.", "s.c.",
                                    "intratumoural"],
                "supplier": ["Janvier", "Charles River", "Envigo", "in-house"],
                "health status": ["SPF", "SOPF", "conventional", "germ-free"],
                "housing": [HOUSING]}
SPECIAL = ("dob", "date of birth", "age", "arrival", "age at arrival")   # their own widgets


def typed_fields(lists: dict | None = None) -> dict:
    """The animal's fields, typed: a list where one is kept, a time for '… time', the
    examples shown while empty."""
    lists = ANIMAL_LISTS if lists is None else lists
    names = list(dict.fromkeys([n for n in lists if n not in ("procedure", "modality",
                                                             "anaesthesia")]
                               + [n for n in FIELD_HINT if n not in SPECIAL]))
    return {n: fdef(n, "list" if lists.get(n) else "time" if timed("", n) else "text",
                    lists.get(n, []), "" if lists.get(n) and n == "housing"
                    else FIELD_HINT.get(n, "")) for n in names}


# what a field is, on hover
FIELD_TIP = {
    "group": "The experimental group the animal belongs to (treated, control…)",
    "supplier": "Where the animal came from (vendor, or bred in-house)",
    "arrival": "The day the animal came into the facility",
    "age at arrival": "The supplier's age on the arrival day — with the arrival date it "
                      "gives the date of birth",
    "protocol": "The ethics / IACUC protocol number the work is done under",
    "housing": "ARRIVE item 15: animals per cage, temperature, light cycle, food and water, "
               "enrichment (nesting material, shelter…); the cage type if it matters",
    "health status": "ARRIVE item 8: the microbiological status the supplier certifies — SPF "
                     "(specific-pathogen-free), SOPF (specific and opportunistic pathogen-free), "
                     "conventional, germ-free",
    "injection route": "How the tracer went in: i.v. tail vein, i.p., s.c.…",
    "euthanasia": "ARRIVE item 9: the method used to kill the animal",
    "euthanasia time": "When the animal was killed: HH:MM, or 2 h p.i.",
    "exclusion reason": "Why this animal's data are left out (ARRIVE item 3)",
    "comment": "Anything else worth reporting",
    "genotype": "Filled in when a listed strain is picked (Options, Strains)",
}

# the other things done to an animal — ARRIVE item 9: one block each, a kind and the fields
# that kind asks for (Options edits them; "+ field" adds one to a single procedure)
EVENT_FIELDS = {"pre/co-injection": ["injection", "what", "when", "dose / amount", "volume",
                                     "route"],
                "imaging": ["modality", "start", "duration"],
                "anaesthesia": ["agent", "induction", "maintenance", "when"],
                "tumour": ["cells", "cell number", "site", "Matrigel", "date"],
                "surgery": ["what", "date", "analgesia"],
                "treatment": ["what", "when", "dose / amount", "route"],
                "diet": ["what", "start date"]}
NEEDS_ANAESTHESIA = ("imaging", "surgery")   # adding one adds an anaesthesia procedure
MODALITIES = ["SPECT", "PET", "CT", "MRI", "planar scintigraphy", "optical"]


def split_modality(events: list[dict]) -> list[dict]:
    """A dual-modality session (SPECT/CT, PET/MRI…) is two imaging procedures, one per
    modality, the same start and duration."""
    out = []
    for ev in events:
        mods = [m.strip() for m in re.split(r"[/+]", str(ev.get("modality", "")))]
        if ev.get("kind") != "imaging" or len(mods) < 2 or not all(mods):
            out.append(ev)
            continue
        out += [dict(ev, modality=m) for m in mods]
    return out
OTHER_EVENT = ["what", "when", "duration"]                  # a kind with no fields of its own


def anaesthesia_apart(events: list[dict]) -> list[dict]:
    """Studies before 2026.10.5.1 kept an imaging session's (a surgery's) anaesthesia as a
    field of it: now a procedure of its own, right after — once for sessions sharing it."""
    out, seen = [], set()
    for ev in events:
        agent = ev.get("anaesthesia", "") if ev.get("kind") in NEEDS_ANAESTHESIA else ""
        out.append({k: v for k, v in ev.items() if not agent or k != "anaesthesia"})
        when = ev.get("start") or ev.get("when") or ev.get("date") or ""
        if agent and (agent, when) not in seen:
            seen.add((agent, when))
            out.append({"kind": "anaesthesia", "agent": agent, "when": when})
    return out


# A field, as Options › Procedures (a kind's fields) and Options › Animal fields describe it:
# how it is typed — text; a list (anything can still be typed); a time on the study day
# (HH:MM, or after the injection: 2 h p.i.); a day (a date, the animal's age then, D-14) —
# the example shown in it while empty, the line it shares with others (a label: "hardware"),
# and when it shows ("modality = SPECT": only when that field holds that).
FIELD_TYPES = {"text": "text", "list": "list", "time": "time · p.i.", "date": "date · age · D-n"}


def fdef(name, type="text", items=(), example="", line="", when="") -> dict:
    return {"name": name, "type": type, "items": list(items), "example": example,
            "line": line, "when": when}


def clean_fdef(d) -> dict | None:
    """A field read back from a file: what is missing takes its default; no name, no field."""
    if not isinstance(d, dict) or not str(d.get("name", "")).strip():
        return None
    out = fdef(str(d["name"]).strip())
    out.update({k: str(d[k]) for k in ("example", "line", "when") if isinstance(d.get(k), str)})
    out["type"] = d.get("type") if d.get("type") in FIELD_TYPES else "text"
    out["items"] = [str(x) for x in d.get("items", []) if str(x).strip()] \
        if isinstance(d.get("items"), list) else []
    return out


_ITEMS = {("imaging", "modality"): MODALITIES,
          ("anaesthesia", "agent"): ["isoflurane", "ketamine / xylazine", "none"],
          ("pre/co-injection", "injection"): ["pre-injection", "co-injection"],
          ("pre/co-injection", "route"): ["i.v. tail vein", "i.p.", "s.c.", "oral gavage"],
          ("pre-injection", "route"): ["i.v. tail vein", "i.p.", "s.c.", "oral gavage"],
          ("diet", "what"): ["standard chow", "iodine-free"]}
_EXAMPLES = {("imaging", "duration"): "30 min", ("anaesthesia", "induction"): "4 % in O2",
             ("anaesthesia", "maintenance"): "1.5-2 % in O2", ("tumour", "cells"): "4T1, TS/A-pc…",
             ("tumour", "cell number"): "1 × 10⁶", ("tumour", "site"): "right flank",
             ("tumour", "Matrigel"): "50 %", ("co-injection", "what"): "blocking dose…",
             ("surgery", "analgesia"): "buprenorphine 0.1 mg/kg"}


def typed_procedures(names: dict[str, list[str]], lists: dict | None = None) -> dict:
    """Procedure kinds given as field names (before 2026.10.6.1) as typed fields: a day or a
    time by the field's name, a list where one was kept for that name."""
    lists = lists or {}
    out = {}
    for kind, fs in names.items():
        out[kind] = []
        for f in fs:
            if kind in NEEDS_ANAESTHESIA and f == "anaesthesia":
                continue                         # a procedure of its own since 2026.10.5.1
            items = _ITEMS.get((kind, f)) or [x for x in lists.get(f, []) if "/" not in x]
            out[kind].append(fdef(f, "date" if dated(kind, f) else "time" if timed(kind, f)
                                  else "list" if items else "text", items,
                                  _EXAMPLES.get((kind, f), "")))
    return out


SPECT_FIELDS = [fdef("system", "list", ["Mediso SPECT/CT"], line="hardware",
                     when="modality = SPECT"),
                fdef("collimator", "list", ["APT62 (mouse, HS)", "APT63 (rat, HS)"],
                     line="hardware", when="modality = SPECT"),
                fdef("isotope / EW", "list", ["99mTc, 140 keV ± 20 %"], line="acquisition",
                     when="modality = SPECT"),
                fdef("time per frame", example="50 s/frame", line="acquisition",
                     when="modality = SPECT")]


def default_procedures() -> dict:
    out = typed_procedures(EVENT_FIELDS)
    out["imaging"] += [dict(f) for f in SPECT_FIELDS]
    return out


def holds(when: str, ev: dict) -> bool:
    """'modality = SPECT, PET': the procedure's modality is one of those (any case); '' —
    always."""
    if "=" not in str(when or ""):
        return True
    f, vals = (x.strip() for x in when.split("=", 1))
    return str(ev.get(f, "")).strip().lower() in {v.strip().lower() for v in vals.split(",")}


def ev_fields(ev: dict, procs: dict[str, list[dict]]) -> list[dict]:
    """The fields a procedure shows: its kind's (those whose `when` holds), any other it
    already holds (as text — what a row holds is never hidden), the note last."""
    own = [f for f in procs.get(ev.get("kind", ""), [fdef(x) for x in OTHER_EVENT])
           if holds(f["when"], ev) or str(ev.get(f["name"], "")).strip()]
    names = {f["name"] for f in own}
    return own + [fdef(k) for k in ev if k not in names and k not in ("kind", "note")] + [
        fdef("note")]


def event_text(ev: dict) -> str:
    """'imaging: modality SPECT, start 13:18, duration 55 min'"""
    return f"{ev.get('kind') or '?'}: " + ", ".join(
        f"{k} {v}" for k, v in ev.items() if k != "kind" and str(v).strip())


# ARRIVE 2.0 items a report checks per animal (item 8: animals, item 9: procedures);
# any of the alternatives fills the item
ARRIVE_NEEDED = [("species", ("species",)), ("strain", ("strain",)), ("sex", ("sex",)),
                 ("age / DOB", ("dob", "date of birth", "age", "age at arrival")),
                 ("weight", ()),
                 ("supplier", ("supplier",)), ("housing", ("housing",)),
                 ("injection route", ("injection route",)),
                 ("euthanasia", ("euthanasia", "euthanasia method"))]
# and per procedure done (item 9): what each kind has to say — anaesthesia goes with the
# procedure that needed it (a biodistribution alone may have none)
EVENT_NEEDED = {"imaging": ["modality", "start"],
                "anaesthesia": ["agent"],
                "surgery": ["what", "date", "analgesia"],
                "tumour": ["cells", "cell number", "site", "date"],
                "treatment": ["what", "dose / amount", "route"],
                "pre/co-injection": ["what", "dose / amount", "route"],
                "pre-injection": ["what", "dose / amount", "route"],     # studies before
                "co-injection": ["what", "dose / amount"],               # 2026.10.6.2
                "diet": ["what", "start date"]}


def dated(kind: str, f: str) -> bool:
    """A procedure field holding a day (not a time on the study day): date / age / D-n."""
    return "date" in f or f == "implanted on" or (
        f == "when" and kind in ("tumour", "surgery", "treatment", "diet"))
# species, strain, genotype (filled in when the strain is picked). Options edits the list.
WT, OUT = "wild type (inbred)", "wild type (outbred)"
STRAINS = [
    ("mouse", "C57BL/6", WT), ("mouse", "C57BL/6J", WT), ("mouse", "C57BL/6N", WT),
    ("mouse", "BALB/c", WT), ("mouse", "BALB/c nude", "Foxn1 nu/nu"),
    ("mouse", "NMRI nude", "Foxn1 nu/nu"), ("mouse", "Athymic nude", "Foxn1 nu/nu"),
    ("mouse", "SCID", "Prkdc scid"), ("mouse", "NSG", "Prkdc scid, Il2rg tm1Wjl"),
    ("mouse", "NOD", WT), ("mouse", "NMRI", OUT), ("mouse", "CD-1", OUT),
    ("mouse", "Swiss", OUT), ("mouse", "FVB", WT), ("mouse", "129", WT),
    ("mouse", "C3H", WT), ("mouse", "DBA/2", WT),
    ("mouse", "ApoE-/- (C57BL/6)", "Apoe tm1Unc"), ("mouse", "Ldlr-/- (C57BL/6)", "Ldlr tm1Her"),
    ("rat", "Wistar", OUT), ("rat", "Sprague-Dawley", OUT), ("rat", "Lewis", WT),
    ("rat", "Fischer 344", WT), ("rat", "RNU nude", "Foxn1 rnu/rnu"),
]


def arrive_gaps(a: "Animal") -> list[str]:
    """The ARRIVE items this animal's card leaves empty, then its procedures' ('imaging:
    anaesthesia')."""
    have = {k.lower() for k, v in a.extra.items() if str(v).strip()}
    return [item for item, keys in ARRIVE_NEEDED
            if (not a.weight_g if item == "weight" else not have & set(keys))] + [
        f"{ev.get('kind')}: {f}" for ev in a.events
        for f in EVENT_NEEDED.get(ev.get("kind", ""), [])
        if not str(ev.get(f, "") or (ev.get("when", "") if f == "start" else "")).strip()] + (
        ["anaesthesia (a procedure, for the imaging or surgery)"]
        if any(e.get("kind") in NEEDS_ANAESTHESIA for e in a.events)
        and not any(e.get("kind") == "anaesthesia" for e in a.events) else [])


_SUP = str.maketrans("⁰¹²³⁴⁵⁶⁷⁸⁹ᵐ", "0123456789m")


def _canon(isotope: str) -> str:
    """'99mTc', 'Tc-99m' and '⁹⁹ᵐTc' all canonicalise to '99mtc'.

    The metastable 'm' is a letter as far as a regex is concerned, so element and mass
    number have to be pulled out together rather than by character class.
    """
    s = re.sub(r"[^0-9a-zA-Z]", "", isotope.translate(_SUP))
    m = re.fullmatch(r"(\d+)(m?)([A-Za-z]+)", s) or re.fullmatch(r"([A-Za-z]+?)(\d+)(m?)", s)
    if not m:
        return s.lower()
    a, b, c = m.groups()
    return (f"{a}{b}{c}" if a.isdigit() else f"{b}{c}{a}").lower()


_CANON_HL = {_canon(k): v for k, v in HALF_LIFE_S.items()}


def set_half_lives(pairs):
    """Replace the isotope table (name, hours) — Options edits it."""
    HALF_LIFE_S.clear()
    HALF_LIFE_S.update({k: float(h) * 3600.0 for k, h in pairs if k and h})
    _CANON_HL.clear()
    _CANON_HL.update({_canon(k): v for k, v in HALF_LIFE_S.items()})


def half_life_s(isotope: str) -> float | None:
    """Half-life in seconds for an isotope written any which way ('99mTc', 'Tc-99m', '⁹⁹ᵐTc')."""
    if not isotope:
        return None
    return HALF_LIFE_S.get(isotope) or _CANON_HL.get(_canon(isotope))


def next_id(last: str, used) -> str:
    """The ID a new animal is offered: 'S1' -> 'S2', '107' -> '108', 'M09' -> 'M10'."""
    m = re.fullmatch(r"(.*?)(\d+)", last or "")
    pre, num, width = (m.group(1), int(m.group(2)), len(m.group(2))) if m else ("", 0, 1)
    while True:
        num += 1
        if (cand := f"{pre}{num:0{width}d}") not in used:
            return cand


def _date(text: str, on: _dt.date) -> _dt.date | None:
    t = parse_time(f"{text.strip()} 00:00", on) if text.strip() else None
    return t.date() if t else None


def parse_dates(text: str, on: _dt.date) -> tuple[list[_dt.date], str]:
    """A date of birth: one date, a range ('7-26/02/26', '07/02-03/03/26', '28/12/25-3/1/26',
    '2026-02-07 – 2026-02-26') or a list ('7/2, 13/2, 26/2/26') — for a cage of animals
    that cannot be told apart. The left side of a range or list borrows what it leaves out
    (month, year) from the last date. Returns (dates, "range" | "list" | "")."""
    s = str(text or "").strip()
    if not s:
        return [], ""
    kind = "list" if re.search(r"[,;]", s) else ""
    if kind:
        parts = re.split(r"\s*[,;]\s*", s)
    else:
        parts = re.split(r"\s*(?:–|—|\bto\b|\s-\s)\s*", s)
        if len(parts) == 1 and s.count("-") == 1 and re.search(r"[/.]", s):
            parts = s.split("-")
        kind = "range" if len(parts) == 2 else ""
    last = _date(parts[-1], on)
    if last is None or len(parts) > 1 and kind == "":
        return [], ""
    out = []
    for x in parts[:-1]:
        n = [int(v) for v in re.findall(r"\d+", x)]
        d = (_date(x, on) if len(n) == 3 else
             _date(f"{n[0]}/{n[1]}/{last.year}", on) if len(n) == 2 else
             _date(f"{n[0]}/{last.month}/{last.year}", on) if len(n) == 1 else None)
        if d is None:
            return [], ""
        if d > last and len(n) < 3:              # 28/12-3/1/26: the year before
            d = d.replace(year=d.year - 1)
        out.append(d)
    return out + [last], kind


def mid_date(text: str, on: _dt.date) -> _dt.date | None:
    """The median of the dates `parse_dates` reads; a range's middle."""
    ds, _ = parse_dates(text, on)
    return _dt.date.fromordinal(int(statistics.median(d.toordinal() for d in ds))) if ds \
        else None


DATE_FORMATS = {"DD/MM/YYYY": "%d/%m/%Y", "DD.MM.YYYY": "%d.%m.%Y", "YYYY-MM-DD": "%Y-%m-%d",
                "DD/MM/YY": "%d/%m/%y"}


def format_date(text: str, on: _dt.date, fmt: str = "DD/MM/YYYY") -> str:
    """A typed date, range or list as `fmt` shows it: '7-26/2/26' -> '07–26/02/2026',
    '7/2-3/3/26' -> '07/02–03/03/2026'. Text that is not a date comes back untouched."""
    ds, kind = parse_dates(text, on)
    f = DATE_FORMATS.get(fmt, "%d/%m/%Y")
    if not ds:
        return str(text or "")
    if kind == "list":
        return ", ".join(d.strftime(f) for d in ds)
    if kind != "range":
        return ds[0].strftime(f)
    a, b = ds
    if f.startswith("%Y"):
        return f"{a:%Y-%m-%d} – {b:%Y-%m-%d}"
    sep = f[2]
    left = (f"{a:%d}" if (a.year, a.month) == (b.year, b.month) else
            f"{a:%d}{sep}{a:%m}" if a.year == b.year else a.strftime(f))
    return f"{left}–{b.strftime(f)}"


_AGE_UNIT = {"d": 1 / 7, "w": 1.0, "m": 30.4375 / 7, "y": 365.25 / 7}


def parse_age(text) -> float | None:
    """An age typed instead of a date of birth, in weeks: '12 wk', '12 weeks', '2.5 mo',
    '84 d', '1 y'; a range '10-12 wk' gives its middle; a bare number is weeks."""
    s = str(text or "").strip().lower().replace(",", ".")
    m = re.fullmatch(r"(\d+(?:\.\d+)?)(?:\s*[-–]\s*(\d+(?:\.\d+)?))?\s*"
                     r"(d|days?|w|wks?|weeks?|m|mo|mos|months?|y|yrs?|years?)?\.?", s)
    if not m:
        return None
    return (float(m[1]) + float(m[2] or m[1])) / 2 * _AGE_UNIT[(m[3] or "w")[0]]


def age_weeks(dob: str, on: _dt.date) -> float | None:
    """Age in weeks on `on` for a date of birth in any of the orders `parse_time` reads
    (2026-06-01, 01.06.2026, 1/6/26 …); a range or a list gives the median's age. An age
    typed instead ('12 wk', '2.5 mo') is taken as it is."""
    d = mid_date(str(dob or ""), on)
    return parse_age(dob) if d is None else (on - d).days / 7


def born(extra: dict, on: _dt.date) -> _dt.date | None:
    """The date of birth a supplier's 'age at arrival' and the arrival date give."""
    got = {k.lower(): v for k, v in extra.items()}
    day, wk = mid_date(str(got.get("arrival", "")), on), parse_age(got.get("age at arrival"))
    return None if day is None or wk is None else day - _dt.timedelta(weeks=wk)


def life_dates(extra: dict, on: _dt.date) -> dict:
    """Date of birth, age on the study day, age at arrival: each from whichever is typed —
    DOB, or the age, or the arrival with the age at arrival (an age typed in the DOB field
    is the age). {"DOB": date, "age": weeks, "age at arrival": weeks}, None where unknown."""
    got = {k.lower(): str(v) for k, v in extra.items() if str(v).strip()}
    typed_dob = got.get("dob", got.get("date of birth", ""))
    dob = mid_date(typed_dob, on)
    age = parse_age(got.get("age", "")) if "age" in got else \
        parse_age(typed_dob) if dob is None else None
    arr, at_arr = mid_date(got.get("arrival", ""), on), parse_age(got.get("age at arrival"))
    if dob is None:
        dob = born(extra, on) or (on - _dt.timedelta(weeks=age) if age is not None else None)
    return {"DOB": dob, "age": (on - dob).days / 7 if dob else age,
            "age at arrival": (arr - dob).days / 7 if arr and dob else at_arr}


def when_views(text: str, dob: _dt.date | None, on: _dt.date) -> dict:
    """A procedure's day, typed one of three ways — a date, the animal's age then ('8 wk'),
    or days from the study day ('D-14') — seen all three ways: {"date", "age", "day"},
    None where it cannot be worked out."""
    s = str(text or "").strip()
    m = re.fullmatch(r"(?i)d\s*([+-−]?)\s*(\d+)", s)
    if m:
        day = on + _dt.timedelta(days=int(m[2]) * (1 if m[1] == "+" else -1))
    elif (day := mid_date(s, on)) is None and (wk := parse_age(s)) is not None and dob:
        day = dob + _dt.timedelta(weeks=wk)
    if not s or day is None:
        return {"date": None, "age": None, "day": None}
    return {"date": day, "age": (day - dob).days / 7 if dob else None, "day": (day - on).days}


def decay(a: float, t_from: _dt.datetime, t_to: _dt.datetime, hl_s: float) -> float:
    """Activity at t_to given activity `a` at t_from."""
    return a * math.exp(-math.log(2) * (t_to - t_from).total_seconds() / hl_s)


def colon_time(s: str) -> str:
    """'12h39' / '12 h 39' / '12H39' -> '12:39' (a French-style time); '2 h p.i.' stays."""
    return re.sub(r"(?i)\b(\d{1,2})\s*h\s*(\d{2})\b", r"\1:\2", str(s or ""))


def parse_time(s: str, day: _dt.date) -> _dt.datetime | None:
    """'16:46' or '16:46:30.25' on `day`, or with a date on either side in any usual order:
    '24/9/26 8:46:37', '24.9 8:47', '8:49 9/24', '2026-09-24 08:49'. Day before month
    unless that makes a month past 12; a date without a year takes `day`'s. A day count
    after the study day, either side: '8:48:21 d1', 'd+1 08:48', '23:10 d-1'.
    ponytail: digits only — '24 Sep' is not understood (reported, so it gets retyped)."""
    if isinstance(s, _dt.datetime):
        return s
    if isinstance(s, _dt.time):
        return _dt.datetime.combine(day, s)
    s = colon_time(s).strip()
    shift = 0
    if dn := re.search(r"(?i)\bd\s*([+-]?\d+)\b(?!:)", s):
        shift, s = int(dn[1]), f"{s[:dn.start()]} {s[dn.end():]}".strip()
    m = re.search(r"(\d{1,2}):(\d{2})(?::(\d{2})(?:[.,](\d+))?)?", s)
    if not m:
        return None
    rest = f"{s[:m.start()]} {s[m.end():]}".strip(" ,T")
    y, mo, d = day.year, day.month, day.day
    if rest:
        n = re.findall(r"\d+", rest)
        if len(n) not in (2, 3) or re.sub(r"[\d\s/.\-]", "", rest):
            return None
        if len(n[0]) == 4:
            y, mo, d = map(int, n)
        else:
            d, mo = int(n[0]), int(n[1])
            y = int(n[2]) if len(n) == 3 else y
            y += 2000 if y < 100 else 0
            if mo > 12:
                d, mo = mo, d
    try:
        return (_dt.datetime(y, mo, d, int(m[1]), int(m[2]), int(m[3] or 0),
                             int((m[4] or "0").ljust(6, "0")[:6]))
                + _dt.timedelta(days=shift))
    except ValueError:
        return None


def parse_pi(text) -> _dt.timedelta | None:
    """A time after the injection: '2 h p.i.', '45 min p.i.', '1h30 p.i.', '1 h 30 min',
    '2.5 h' (a unit or 'p.i.' needed: a bare number is not read)."""
    s = str(text or "").strip().lower().replace(",", ".")
    s, pi = re.subn(r"\s*(?:p\.?\s*i\.?|post[- ]?inj\w*\.?)$", "", s)
    m = re.fullmatch(r"(?:(\d+(?:\.\d+)?)\s*h(?:ours?|rs?)?)?\s*"
                     r"(?:(\d+(?:\.\d+)?)\s*(min(?:utes?|s)?|m|')?)?", s.strip())
    if not m or not (m[1] or m[2]) or m[2] and not m[1] and not m[3] and not pi:
        return None
    return _dt.timedelta(hours=float(m[1] or 0), minutes=float(m[2] or 0))


def format_pi(d: _dt.timedelta) -> str:
    """'2 h p.i.', '1 h 30 min p.i.', '45 min p.i.'; before the injection: '20 min before'."""
    mins = round(abs(d.total_seconds()) / 60)
    h, m = divmod(mins, 60)
    text = " ".join([f"{h} h"] * bool(h) + [f"{m} min"] * bool(m or not h))
    return text + (" before" if d.total_seconds() < 0 else " p.i.")


def time_views(text: str, inj: _dt.datetime | None, day: _dt.date) -> dict:
    """A time on the study day typed as a clock time ('13:10') or after the injection ('2 h
    p.i.'), seen both ways: {"time": datetime, "pi": timedelta}, None where unknown."""
    pi = parse_pi(text)
    if pi is not None:
        return {"time": inj + pi if inj else None, "pi": pi}
    t = parse_time(str(text or ""), day) if str(text or "").strip() else None
    return {"time": t, "pi": t - inj if t and inj else None}


def timed(kind: str, f: str) -> bool:
    """A field holding a time on the study day (a clock time or 'x h p.i.'), by its name:
    '… time', 'start', 'when' (but a day-kind's 'when' is a date) — how fields named so
    before 2026.10.6.1 are typed; Options › Procedures says it since."""
    return not dated(kind, f) and ("time" in f.lower() or f in ("start", "when"))


def event_time(a: "Animal", text: str, day: _dt.date) -> _dt.datetime | None:
    """A procedure's (a field's) time, however typed: a clock time or after this animal's
    injection."""
    return time_views(text, parse_time(a.inj_time, day) if a.inj_time else None, day)["time"]


def time_text(a: "Animal", text: str, day: _dt.date) -> str:
    """'14:13 (2 h 35 min p.i.)': a typed time with the other way of saying it, for a report."""
    v = time_views(text, parse_time(a.inj_time, day) if a.inj_time else None, day)
    return f"{v['time']:%H:%M} ({format_pi(v['pi'])})" if v["time"] and v["pi"] is not None \
        else str(text or "")


ANIMAL_SPECS = typed_fields()   # how the animal window offers each field (fixed)


TIME_FORMATS = {"HH:MM d+n": "08:48 d+1", "HH:MM:SS d+n": "08:48:21 d+1",
                "DD/MM HH:MM": "25/09 08:48", "YYYY-MM-DD HH:MM": "2026-09-25 08:48"}


def format_time(s: str, day: _dt.date, fmt: str = "HH:MM d+n") -> str:
    """A typed time as `fmt` shows it ('HH:MM d+n': '08:48' on the study day, '08:48 d+1'
    the day after); text that is not a time comes back untouched."""
    t = parse_time(s, day)
    if t is None:
        return str(s or "")
    if fmt.endswith("d+n"):
        n = (t.date() - day).days
        return t.strftime("%H:%M:%S" if "SS" in fmt else "%H:%M") + (f" d{n:+d}" if n else "")
    return t.strftime("%d/%m %H:%M" if fmt.startswith("DD") else "%Y-%m-%d %H:%M")


def time_problems(a: "Animal", day: _dt.date) -> dict[str, str]:
    """Card times that cannot be right, field -> why: the full syringe is read before the
    injection, the empty one and the injection site after it."""
    t = {f: parse_time(getattr(a, f), day) for f in ("full_time", "inj_time", "empty_time",
                                                      "tail_time")}
    out = {}
    if t["full_time"] and t["empty_time"] and t["empty_time"] < t["full_time"]:
        out["empty_time"] = "empty syringe read before the full one"
    if inj := t["inj_time"]:
        if t["full_time"] and t["full_time"] > inj:
            out["full_time"] = "full syringe read after the injection"
        if t["empty_time"] and t["empty_time"] < inj:
            out["empty_time"] = "empty syringe read before the injection"
        if t["tail_time"] and t["tail_time"] < inj:
            out["tail_time"] = "injection site read before the injection"
    return out


DT_WARN = 1.1            # dead-time factor above which a counting is out of the target range
DT_BEFORE = 1.5          # ... in a study saved before 2026.10.6.2 that does not say
DT_VALID = 1.5           # ... past which it is not valid: flagged, used only if no other is
MIN_COUNTS = 10000       # a counting aimed at (±1 %); fewer is still taken when it is all there is
VALID_COUNTS = 1000      # ... under which it is not valid (±3 %)
BLANK_COUNTS = 1000      # a blank vial counting this many is not empty (background: 100-600)
RECOUNT_COUNTS = 1000    # vials compared to tell a recount (±3 %, under its 10 % test)
DRIFT_MODES = {"": "off", "scale": "scale: each tube by the control tubes' ratio",
               "offset": "offset: each tube minus the control tubes' change"}


# --------------------------------------------------------------------------- model
@dataclass
class Animal:
    id: str = ""
    aliases: list[str] = field(default_factory=list)
    weight_g: float | None = None
    isotope: str = "99mTc"
    molecule: str = ""
    full_mbq: float | None = None
    full_time: str = ""
    empty_mbq: float | None = None
    empty_time: str = ""
    inj_time: str = ""
    tail_mbq: float | None = None                # injection site read by hand: wins over a
    tail_time: str = ""                          # counted "tail" vial
    losses: list[dict[str, str]] = field(default_factory=list)   # other activity that did
    #             not go in, read by hand (a cotton held on the tail…): mbq (as typed), time, what
    bio_notes: list[str] = field(default_factory=list)   # notes on the biodistribution
    note: str = ""
    extra: dict[str, str] = field(default_factory=lambda: {"species": "mouse"})
    typed: dict[str, str] = field(default_factory=dict)   # numbers as read: "1.30", not 1.3
    show: dict[str, bool] = field(default_factory=dict)   # card fields shown / hidden by hand
    events: list[dict[str, str]] = field(default_factory=list)   # kind + its fields -> text

    @property
    def label(self) -> str:
        names = [x for x in self.aliases if x]
        return f"{self.id} ({', '.join(names)})" if names else self.id

    @staticmethod
    def from_dict(d: dict) -> "Animal":
        """Studies saved before aliases became a list carry a single `alias` string."""
        d = dict(d)
        old = d.pop("alias", "")
        d.setdefault("aliases", [old] if old else [])
        mbq, t = d.pop("loss_mbq", None), d.pop("loss_time", "")   # 2026.10.5: one loss
        if mbq is not None or t:
            d.setdefault("losses", []).append(
                {"mbq": d.get("typed", {}).pop("loss_mbq", "" if mbq is None else f"{mbq:g}"),
                 "time": t, "what": ""})
        d["events"] = split_modality(anaesthesia_apart(d.get("events", [])))
        ex = d.get("extra", {})
        for f, kind, key in (("tumour", "tumour", "cells"),       # 2026.10.6: procedures
                             ("anaesthesia", "anaesthesia", "agent"), ("diet", "diet", "what")):
            v = str(ex.pop(f, "") or "").strip()
            ev = next((e for e in d["events"] if e.get("kind") == kind), None)
            if not v or ev and v in (ev.get(key), ev.get("note")):
                continue
            if ev is None:
                d["events"].append({"kind": kind, key: v})
            elif not ev.get(key):
                ev[key] = v
            else:                                # both said something: neither is lost
                ev["note"] = "; ".join(x for x in (ev.get("note", ""), f"{f}: {v}") if x)
        return Animal(**{k: v for k, v in d.items() if k in Animal.__dataclass_fields__})


@dataclass
class Tissue:
    name: str = ""
    role: str = "tissue"
    hand: list[str] = field(default_factory=list)  # "mass" / "activity" / "note": typed rows
    closed: list[str] = field(default_factory=list)  # typed rows closed: kept, not used
    batch: str = ""              # the runs it is counted in: "" = the main ones, else a name


@dataclass
class Source:
    """A Hidex file placed onto the grid."""

    uid: str = ""
    path: str = ""
    kind: str = "count"          # one of KINDS
    preset: str = ANIMAL_MAJOR
    animals: list[str] = field(default_factory=list)
    tissues: list[str] = field(default_factory=list)
    slotmap: dict[str, list[str]] = field(default_factory=dict)  # "rack:vial" -> [animal, tissue]
    batch: str = ""              # the tissue run it was guessed into
    auto: bool = True            # kind and mapping are guessed until the user sets them
    stamp: dict = field(default_factory=dict)    # the file as it was used: size, times, sha256

    @staticmethod
    def from_dict(d: dict) -> "Source":
        """Studies saved before kinds became empty/filled carry kind=tare plus a role."""
        d = {k: v for k, v in d.items() if k in Source.__dataclass_fields__ or k == "role"}
        role = d.pop("role", "")
        if d.get("kind") == "tare":
            d["kind"] = role or "empty"
        d.setdefault("auto", False)
        return Source(**d)

    def pairs(self) -> list[tuple[str, str]]:
        if self.preset == TISSUE_MAJOR:
            return [(a, t) for t in self.tissues for a in self.animals]
        return [(a, t) for a in self.animals for t in self.tissues]

    def mapping(self, run: hidex.Run) -> dict[str, tuple[str, str]]:
        """slot key -> (animal, tissue). The preset fills it; slotmap overrides per slot, and
        an empty (or half-typed) slotmap entry takes the slot out."""
        out = {s.key: p for s, p in zip(run.slots, self.pairs())}
        for k, v in self.slotmap.items():
            if len(v) > 1 and v[0] and v[1]:
                out[k] = (v[0], v[1])
            else:
                out.pop(k, None)
        return out


@dataclass
class Manual:
    """A hand-entered mass and/or dose-calibrator activity for one cell."""

    animal: str = ""
    tissue: str = ""
    mass_g: float | None = None
    mbq: float | None = None
    time: str = ""
    note: str = ""
    typed: dict[str, str] = field(default_factory=dict)   # numbers as read


@dataclass
class Cell:
    mass_g: float | None = None
    bq: float | None = None                      # at the study reference time
    mass_src: str = ""
    bq_src: str = ""
    mass_alts: list[tuple[str, float]] = field(default_factory=list)   # src, g
    empties: list[tuple[str, float]] = field(default_factory=list)     # empty-tube src, g
    fulls: list[tuple[str, float, str]] = field(default_factory=list)  # src, g, filled |
    #                                  total (a weigh+count run's) | direct (a mass as it is)
    mass_empty: str = ""                         # the empty tube the mass came off
    windows: list[tuple[str, str, float]] = field(default_factory=list)  # src, window, Bq
    dead_time: float | None = None
    counts: float | None = None
    raw: dict[str, tuple] = field(default_factory=dict)   # count file -> (counts, cpm) as read
    alts: list[tuple] = field(default_factory=list)   # src, bq, dead time, counts (or cpm),
    #                                              time, relative σ (`rsd`)
    flags: list[str] = field(default_factory=list)
    bq_used: list[str] = field(default_factory=list)     # the countings (files) behind `bq`
    mass_used: list[str] = field(default_factory=list)   # the weighings behind `mass_g`
    calib: float | None = None                   # a dose-calibrator reading typed, at `ref`
    rule_w: dict[str, str] = field(default_factory=dict)  # count file -> the rule's window
    every: list[tuple] = field(default_factory=list)     # every counting in every window:
    #             src, bq, dead time, counts (or cpm), time, window, cpm as read, the rule's?,
    #             counts as counted, relative σ
    pairs: dict[tuple, list[str]] = field(default_factory=dict)   # ("count", file, window) |
    #             ("mass", file) | ("tare", file) -> how it differs from each other one

    def spread_pct(self, study: "Study") -> float | None:
        """Largest disagreement between the valid repeat counts of this vial, in percent."""
        vals = [x[1] for x in self.alts if x[1] > 0 and is_valid(x, study)]
        if len(vals) < 2:
            return None
        return (max(vals) - min(vals)) / min(vals) * 100.0


@dataclass
class Study:
    name: str = ""
    date: str = ""                               # ISO experiment date
    animals: list[Animal] = field(default_factory=list)
    tissues: list[Tissue] = field(default_factory=list)
    sources: list[Source] = field(default_factory=list)
    manual: list[Manual] = field(default_factory=list)
    window: str = ""                             # short window label; "" = by window_rule
    window_rule: str = "wide"                    # one isotope: widest window, else its peak
    min_counts: float = MIN_COUNTS               # the target range: this many counts, …
    min_basis: str = "counts"                    # ... counts or cpm (the validity's too)
    valid_counts: float = VALID_COUNTS           # valid: this many at least, and a dead time …
    valid_dt: float = DT_VALID                   # … ≤ this — else flagged, and used only when
    #                                              no counting of the vial is valid
    pick_count: str = "first"                    # first | last | auto (most counts) | all
    #             (combined) | round:<file> — of the countings in the target range (and that
    #             agree, if count_agree); none in range: of the valid ones; none valid: the
    #             nearest to valid
    count_agree: bool = True                     # a valid counting out of the consensus
    #                                              (`agreeing`) is left out before the rule
    count_tol_pct: float = 3.0                   # two countings agree within this % …
    count_tol_sigma: float = 3.0                 # … or this many σ of their counting statistics
    dt_max: float = DT_WARN                      # … dead time ≤ this,
    cpm_max: float = 0                           # CPM ≤ this (0: no bound)
    combine: str = "weighted"                    # several countings: weighted (by counts) | mean
    pick_mass: str = "auto"                      # auto (typed wins) | files
    mass_rule: str = "first"                     # a tube weighed more than once: first (the
    #             day-of weighing) | last | mean | median — of those that agree, if mass_agree
    mass_agree: bool = True                      # a weighing out of the consensus is left
    #                                              out before the rule
    mass_tol_mg: float = 2.0                     # two weighings agree within this many mg …
    mass_tol_pct: float = 10.0                   # … or this % of the tissue
    drift_fix: str = ""                          # weighings corrected by their control tubes:
    #                                              "" | "scale" | "offset" (DRIFT_MODES)
    batches: dict[str, str] = field(default_factory=dict)   # tissue run -> vial order
    subtract_tail: bool = True
    chosen: list[list[str]] = field(default_factory=list)   # [animal, tissue, "count" | "mass"
    #             | "empty", file, file…] picked in the Results for one cell: several countings
    #             combined as `combine` says, several weighings give their mean; a counting
    #             in another window than the rule's is 'file@112-168' ("dose calibrator",
    #             "manual" too)
    ref_time: str = ""                           # ISO; "" = derive from the counting files
    ref_rule: str = "injection"                  # Bq shown at each animal's injection time |
    #                                              "time": at ref_time (typed: 18:00, 1/10 9:00;
    #                                              empty: the first counting's reference)
    #                                              "study": at the one instant above
    output: list[str] = field(default_factory=list)   # tissues the results and report show,
    #                                              in order; [] = every tissue, as in the grid
    animal_output: list[str] = field(default_factory=list)   # the same for the animals (ids)
    group_by: str = ""                           # "molecule": the animals grouped by it
    tissue_labels: dict[str, str] = field(default_factory=dict)   # name -> as the results and
    animal_labels: dict[str, str] = field(default_factory=dict)   # id -> the report show it
    file_eff: bool = True                        # the counter files' own efficiencies (Hidex)
    efficiency: dict[str, list] = field(default_factory=dict)   # "counter|window" -> [isotope,
    #             energy window, counts per decay]: typed, for files with none (Wizard2, a
    #             Hidex set to 1) or all of them when file_eff is off
    ranges: list[list] = field(default_factory=list)   # [tissue, mg min, mg max, %IA/g min,
    #                                              max] expected; 0 = no bound — flags only
    embedded: dict[str, dict] = field(default_factory=dict)  # path -> the file's data, packed
    outside_edit = False                         # loaded: changed since saved (not a field)

    # ------------------------------------------------------------------ persistence
    def to_json(self) -> str:
        """Readable: a short list on one line; the files' data kept compact at the end, and
        a fingerprint of the whole, to tell an edit made outside BioDist."""
        d = json.loads(json.dumps({"biodist": 1, **asdict(self)}, default=str))
        packed = d.pop("embedded")
        text = json.dumps(d, indent=1, ensure_ascii=False)
        text = re.sub(r"\[\s*\n\s*([^\[\]{}]*?)\s*\]",       # [\n "107",\n "Liver"\n] -> one line
                      lambda m: "[" + re.sub(r",\s*\n\s*", ", ", m[1]) + "]", text)
        names = list(packed)
        d["embedded"] = {"files": names, "data": [packed[k] for k in names]}
        files = ",\n".join(f"   {json.dumps(k, ensure_ascii=False)}" for k in names)
        data = ",\n".join(f"   {json.dumps(packed[k], separators=(',', ':'))}" for k in names)
        tail = (',\n "embedded": {\n  "files": [' + (f"\n{files}\n  " if names else "")
                + '],\n  "data": [' + (f"\n{data}\n  " if names else "") + ']\n },\n'
                f' "saved_sha256": "{_fingerprint(d)}"\n}}')
        return text[:-2] + tail

    @staticmethod
    def from_json(text: str) -> "Study":
        d = json.loads(text)
        sha = d.pop("saved_sha256", None)
        edited = sha is not None and sha != _fingerprint(d)
        d.pop("biodist", None)
        emb = d.get("embedded") or {}
        if "files" in emb and "data" in emb:         # since 2026.10.5.3: a list, then the data
            d["embedded"] = dict(zip(emb["files"], emb["data"]))
        if d.get("average") in ("mean", "median"):   # 2026.10.5: one rule for both
            d["mass_rule"] = d["average"]
            if d.get("pick_count", "auto") == "auto":
                d["pick_count"] = d["average"]
        d.setdefault("dt_max", DT_BEFORE)            # what a study saved before held to
        d.setdefault("valid_dt", d["dt_max"])        # one range until 2026.10.7: its flags
        d.setdefault("valid_counts", d.get("min_counts", MIN_COUNTS))
        d.setdefault("pick_count", "auto")
        d.setdefault("mass_rule", "first")
        d.setdefault("count_agree", False)           # a study saved before took every one
        d.setdefault("mass_agree", False)
        if d.get("pick_count") in ("mean", "median", "pooled"):   # 2026.10.6.1: "all",
            d.setdefault("combine", "weighted" if d["pick_count"] == "pooled" else "mean")
            d["pick_count"] = "all"                  # combined as `combine` says
        s = Study(**{k: v for k, v in d.items() if k in Study.__dataclass_fields__
                     and k not in ("animals", "tissues", "sources", "manual")})
        s.animals = [Animal.from_dict(a) for a in d.get("animals", [])]
        s.tissues = [Tissue(**t) for t in d.get("tissues", [])]
        for t, n in zip(s.tissues, unique_names([t.name for t in s.tissues])):
            t.name = n                           # files from before 2026.9.30.1 may repeat one
        s.sources = [Source.from_dict(x) for x in d.get("sources", [])]
        s.manual = [Manual(**m) for m in d.get("manual", [])]
        s.outside_edit = edited
        return s

    def save(self, path):
        Path(path).write_text(self.to_json(), encoding="utf-8")

    @staticmethod
    def load(path) -> "Study":
        return Study.from_json(Path(path).read_text(encoding="utf-8"))

    # ---------------------------------------------------------------------- helpers
    @property
    def day(self) -> _dt.date:
        try:
            return _dt.date.fromisoformat(self.date)
        except ValueError:
            return _dt.date.today()

    def animal(self, aid: str) -> Animal | None:
        return next((a for a in self.animals if a.id == aid), None)

    def tissue(self, name: str) -> Tissue | None:
        return next((t for t in self.tissues if t.name == name), None)

    def source(self, uid: str) -> Source | None:
        return next((s for s in self.sources if s.uid == uid), None)

    def output_tissues(self) -> list[Tissue]:
        """What the results and the report list: the output list, else every collected
        tissue (role 'tissue') in the grid's order."""
        if self.output:
            return [t for n in self.output if (t := self.tissue(n))]
        return [t for t in self.tissues if t.role == "tissue"]

    def output_animals(self) -> list[Animal]:
        """The animals the results and the report show, in their order (else the cards'),
        grouped by molecule if asked — each group where its first animal is."""
        out = [a for i in self.animal_output if (a := self.animal(i))] if self.animal_output \
            else list(self.animals)
        if self.group_by == "molecule":
            first = {m: k for k, m in reversed(list(enumerate(a.molecule for a in out)))}
            out.sort(key=lambda a: first[a.molecule])
        return out

    def column_label(self, a: Animal) -> str:
        lab = self.animal_labels.get(a.id, a.label)
        return f"{lab}\n{a.molecule}" if self.group_by == "molecule" else lab

    def tissue_label(self, name: str) -> str:
        return self.tissue_labels.get(name, name)

    def rename_tissue(self, old: str, new: str):
        """A rename follows the tissue everywhere it is named, and re-guesses an unset role."""
        t = self.tissue(old)
        if not t or not new or new == old:
            return
        if t.role == guess_role(old):
            t.role = guess_role(new)
        t.name = new
        for m in self.manual:
            if m.tissue == old:
                m.tissue = new
        for x in self.chosen:
            if x[1] == old:
                x[1] = new
        self.output = [new if x == old else x for x in self.output]
        if old in self.tissue_labels:
            self.tissue_labels[new] = self.tissue_labels.pop(old)
        for s in self.sources:
            s.tissues = [new if x == old else x for x in s.tissues]
            for v in s.slotmap.values():
                if v[1:2] == [old]:
                    v[1] = new

    def animal_renamed(self, old: str, new: str):
        """An animal's ID changed (already on the animal): everything that names it follows —
        the files' placements, the values typed for it, its picks, its label, its place in
        the output."""
        for m in self.manual:
            if m.animal == old:
                m.animal = new
        for x in self.chosen:
            if x[0] == old:
                x[0] = new
        self.animal_output = [new if x == old else x for x in self.animal_output]
        if old in self.animal_labels:
            self.animal_labels[new] = self.animal_labels.pop(old)
        for s in self.sources:
            s.animals = [new if x == old else x for x in s.animals]
            for v in s.slotmap.values():
                if v[:1] == [old]:
                    v[0] = new

    def batch_names(self) -> list[str]:
        """The tissue runs, main ("") first, then in list order."""
        return list(dict.fromkeys([""] + [t.batch for t in self.tissues]))

    def batch_tissues(self, batch: str) -> list[str]:
        return [t.name for t in self.tissues if t.batch == batch]

    def batch_order(self, batch: str) -> str:
        """Main runs go animal by animal; the others, collected on the spot, tissue by tissue."""
        return self.batches.get(batch, TISSUE_MAJOR if batch else ANIMAL_MAJOR)

    def new_uid(self) -> str:
        n = 1
        while any(s.uid == f"s{n}" for s in self.sources):
            n += 1
        return f"s{n}"

    def hl_for(self, aid: str) -> float:
        a = self.animal(aid)
        return (half_life_s(a.isotope) if a else None) or HALF_LIFE_S["99mTc"]

    def window_for(self, run: hidex.Run) -> str | None:
        """The full window name in `run` for the study's window choice; with none, by rule:
        a study with one isotope takes its widest window (all the counts), several take each
        one's photopeak (the narrowest) so they do not spill into each other.
        ponytail: one window per run — a run mixing isotopes gets the first one's."""
        if not run.windows:
            return None
        if self.window:
            return next((w for w, short in zip(run.windows, run.short_windows())
                         if short == self.window), run.windows[0])
        isos = {_canon(a.isotope) for a in self.animals if a.isotope}
        own = [w for w in run.windows if _canon(w.split("_")[0]) in isos] or run.windows

        def width(w):
            n = [float(x) for x in re.findall(r"\d+(?:\.\d+)?", w.split("_", 1)[-1])]
            return n[-1] - n[0] if len(n) >= 2 else 0
        wide = len(isos) <= 1 and self.window_rule == "wide"
        return (max if wide else min)(own, key=width)

    def eff_for(self, run: hidex.Run, w: str) -> tuple[float | None, bool]:
        """(counts per decay, from the file?) for a window of a file: the file's own unless
        the study says otherwise or it has none (1 = none set in the counter)."""
        own = run.efficiency.get(w)
        own = own if own and abs(own - 1) > 1e-9 else None
        typed = _float((self.efficiency.get(f"{run.counter}|{w}") or [None] * 3)[2])
        if self.file_eff and own:
            return own, True
        return (typed, False) if typed else (None, False)

    def eff_rows(self, runs) -> list[tuple[str, str, str, float | None]]:
        """The counter·window pairs the counting files hold: (key, counter, window, the
        files' efficiency or None) — what the efficiency table lists."""
        out = {}
        for s in self.sources:
            r = runs.get(s.path)
            if r and s.kind in ("count", "weigh_count") and r.windows:
                w = self.window_for(r)
                e = r.efficiency.get(w)
                out.setdefault(f"{r.counter}|{w}", (r.counter, w, e))
        return [(k, *v) for k, v in out.items()]

    def windows_available(self, runs) -> list[str]:
        out = []
        for r in runs.values():
            for s in r.short_windows():
                if s not in out:
                    out.append(s)
        return out

    def reference(self, runs) -> _dt.datetime:
        """One instant every activity is corrected to."""
        if self.ref_time:
            got = parse_time(self.ref_time, self.day)
            if got:
                return got
        norms = sorted(r.normalized_to for r in runs.values() if r.normalized_to)
        if norms:
            return norms[0]
        times = sorted(s.time for r in runs.values() for s in r.slots if s.time)
        return times[0] if times else _dt.datetime.combine(self.day, _dt.time(12, 0))


# ------------------------------------------------------------------------ computing
@dataclass
class Result:
    ref: _dt.datetime
    cells: dict[tuple[str, str], Cell]
    injected_bq: dict[str, float]                # animal id -> injected activity at ref, tail already removed
    tail_bq: dict[str, float]
    notes: list[str] = field(default_factory=list)
    rounds: list[list[str]] = field(default_factory=list)   # count files, one pass over the vials
    spans: dict[str, tuple] = field(default_factory=dict)   # file -> (first, last) vial time
    suggest: list[dict] = field(default_factory=list)   # findings with a fix to offer
    refs: dict[str, _dt.datetime] = field(default_factory=dict)   # animal -> the instant its
    #                                    Bq are shown at (its injection, or `ref`)
    hls: dict[str, float] = field(default_factory=dict)          # animal -> half-life (s)
    files: dict[str, tuple] = field(default_factory=dict)   # file -> (kind, started)
    no_eff: list[str] = field(default_factory=list)  # counter|window with no efficiency

    def ref_of(self, aid: str) -> _dt.datetime:
        return self.refs.get(aid, self.ref)

    def at_ref(self, aid: str, bq: float | None) -> float | None:
        """An activity held at `ref` (where every sum is made), at the animal's own instant."""
        if bq is None or aid not in self.refs:
            return bq
        return decay(bq, self.ref, self.refs[aid], self.hls[aid])

    def round_of(self, src: str) -> int | None:
        return next((i for i, r in enumerate(self.rounds) if src in r), None)

    def round_label(self, i: int) -> str:
        """'round 2 · 24 Sep 15:30–17:01 · 6 files'"""
        t = [x for n in self.rounds[i] for x in self.spans.get(n, ()) if x]
        when = (f"{min(t):%d %b %H:%M}–{max(t):%H:%M}" if min(t).date() == max(t).date()
                else f"{min(t):%d %b %H:%M} – {max(t):%d %b %H:%M}") if t else ""
        n = len(self.rounds[i])
        return f"round {i + 1} · {when} · {n} file{'s' if n > 1 else ''}"

    def cell(self, aid: str, tissue: str) -> Cell:
        return self.cells.get((aid, tissue), Cell())

    def weighing_label(self, name: str, kind: str) -> str:
        """'weight', 'count+weight r2', 'typed': a weighing named by the pass it came in."""
        i = self.round_of(name)
        return ("typed" if name == "manual" else "weight" if kind == "filled" else
                f"count+weight r{i + 1}" if i is not None else KIND_LABEL.get(
                    self.files.get(name, ("",))[0], "weighed"))

    def source_label(self, aid: str, tissue: str, what: str) -> str:
        """Where a cell's activity or mass comes from, in a word or three — the Results'
        'activity source' / 'mass source': 'round 1', 'round 2 · 112-168', 'calibrator',
        'mean r1+r2'; 'weight − tare', 'count+weight r2 − tare', 'mean of weight, r2 − tare',
        'typed'."""
        c = self.cell(aid, tissue)
        if what == "activity":
            if c.bq_src == "dose calibrator":
                return "calibrator"
            rs = [(f"r{i + 1}" if (i := self.round_of(n.rsplit("@", 1)[0])) is not None
                   else _stem(n.rsplit("@", 1)[0])) + (f" · {n.rsplit('@', 1)[1]}" if "@" in n
                                                       else "") for n in c.bq_used]
            return "" if not rs else f"round {rs[0][1:]}" if len(rs) == 1 and \
                rs[0][0] == "r" else c.bq_src.split(" of ")[0] + " " + "+".join(rs)
        kinds = {n: k for n, _, k in c.fulls}
        labs = [self.weighing_label(n, kinds.get(n, "")) for n in c.mass_used if n in kinds]
        if not labs:
            return ""
        tare = " − tare" if any(kinds[n] in ("filled", "total") for n in c.mass_used
                                if n in kinds) else ""
        if len(labs) == 1:
            return labs[0] + tare
        return f"{c.mass_src.split(' of ')[0]} of " + ", ".join(
            x.replace("count+weight ", "") for x in labs) + tare

    def value(self, study: Study, aid: str, tissue: str, unit: str, bq: float | None = None,
              mass: float | None = None) -> float | None:
        """One number for the results grid. Returns None when the inputs are not there.
        `bq` / `mass` given: what the cell would read with them instead."""
        c = self.cell(aid, tissue)
        if bq is not None or mass is not None:
            c = Cell(mass_g=c.mass_g if mass is None else mass, bq=c.bq if bq is None else bq,
                     raw=c.raw, bq_src=c.bq_src, dead_time=c.dead_time)
        idb = self.injected_bq.get(aid)
        if unit == "mass":
            return c.mass_g
        if unit == "mass_mg":
            return None if c.mass_g is None else c.mass_g * 1000
        if unit == "bq":
            return self.at_ref(aid, c.bq)
        if unit in ("counts", "cpm"):            # the counting in use, as the counter read it
            return c.raw.get(c.bq_src, (None, None))[unit == "cpm"]
        if unit == "dt":
            return c.dead_time
        if c.bq is None:
            return None
        if unit == "bq_g":
            return self.at_ref(aid, c.bq) / c.mass_g if c.mass_g else None
        if idb is None or idb <= 0:
            return None
        pid = c.bq / idb * 100.0
        if unit == "pid":
            return pid
        if unit == "pid_g":
            return pid / c.mass_g if c.mass_g else None
        if unit == "suv":
            a = study.animal(aid)
            if not a or not a.weight_g or not c.mass_g:
                return None
            return (c.bq / c.mass_g) / (idb / a.weight_g)
        return None


UNITS = [("pid_g", "%IA/g"), ("pid", "%IA"), ("suv", "SUV"),
         ("bq_g", "Bq/g"), ("bq", "Bq"), ("mass", "mass (g)"), ("mass_mg", "mass (mg)"),
         ("counts", "counts"), ("cpm", "CPM"), ("dt", "dead time")]
DIGITS = {"pid_g": 2, "pid": 2, "suv": 2, "bq_g": 0, "bq": 0, "mass": 4, "mass_mg": 1,
          "mbq": 4, "counts": 0, "cpm": 0, "dt": 3}     # decimals shown, per unit


def _fingerprint(d: dict) -> str:
    return hashlib.sha256(json.dumps(d, sort_keys=True, separators=(",", ":"),
                                     ensure_ascii=False).encode()).hexdigest()


def stamp(path) -> dict:
    """A file as it is now: size, times, SHA-256 — to tell later whether it changed."""
    p = Path(path)
    st = p.stat()
    return {"size": st.st_size,
            "modified": _dt.datetime.fromtimestamp(st.st_mtime).isoformat(timespec="seconds"),
            "created": _dt.datetime.fromtimestamp(st.st_ctime).isoformat(timespec="seconds"),
            "sha256": hashlib.sha256(p.read_bytes()).hexdigest()}


def load_runs(study: Study) -> dict[str, hidex.Run]:
    """Read every source file once; a file gone, the copy the study keeps. Missing/broken
    files are simply absent."""
    runs = {}
    for s in study.sources:
        if s.path and s.path not in runs:
            try:
                runs[s.path] = hidex.read(s.path)
            except Exception:  # noqa: BLE001 — reported by compute() as a note
                if s.path in study.embedded:
                    runs[s.path] = hidex.unpack(s.path, study.embedded[s.path])
    return runs


def compute(study: Study, runs: dict[str, hidex.Run] | None = None) -> Result:
    """Fold every source into the tissue x animal grid and correct it all to one instant."""
    runs = load_runs(study) if runs is None else runs
    live = [s for s in study.sources if s.kind != "ignored"]     # kept on screen, not used
    runs = {s.path: runs[s.path] for s in live if s.path in runs}
    ref = study.reference(runs)
    cells: dict[tuple[str, str], Cell] = {}
    notes: list[str] = []

    for s in live:
        if s.path and s.path not in runs:
            notes.append(f"could not read {Path(s.path).name}")

    def cell(key) -> Cell:
        return cells.setdefault(key, Cell())

    # ---- weighings, cell by cell: empty tubes, filled tubes (or a weigh+count run's total),
    # and masses as they are (a weigh+count run that tared itself, a mass typed by hand)
    for s in live:
        run = runs.get(s.path)
        if not run or s.kind == "count":
            continue
        name, mapping = Path(s.path).name, s.mapping(run)
        for slot in run.slots:
            key = mapping.get(slot.key)
            if not key or not key[0]:
                continue
            if slot.sample_g is not None:
                cell(key).fulls.append((name, slot.sample_g, "direct"))
            elif s.kind == "empty" and slot.tare_g is not None:
                cell(key).empties.append((name, slot.tare_g))
            elif (g := slot.tare_g if s.kind == "filled" else slot.total_g) is not None:
                cell(key).fulls.append((name, g, "filled" if s.kind == "filled" else "total"))
    for m in study.manual:
        t = study.tissue(m.tissue)
        if m.animal and m.tissue and m.mass_g is not None and not (t and "mass" in t.closed):
            cell((m.animal, m.tissue)).fulls.append(("manual", m.mass_g, "direct"))

    # ---- the control tubes (blanks), weighed empty then filled, should not change: a
    # balance that reads them light reads every tube of that file light, in proportion
    # (260930: 0.02-0.03 %, 1-2 mg on a 6 g tube, the same on the 2.5 g bare vials)
    blanks = {t.name for t in study.tissues if t.role == "blank"}
    ratios: dict[str, list[float]] = {}
    shifts: dict[str, list[float]] = {}
    for (aid, tn), c in cells.items():
        if tn not in blanks or not c.empties:
            continue
        e = c.empties[-1][1]
        for name, g, kind in c.fulls:
            if kind == "filled" and abs(g - e) < 0.005:     # beyond that: not the same tube
                ratios.setdefault(name, []).append(g / e)
                shifts.setdefault(name, []).append(g - e)
            elif kind == "filled" and g - e >= 0.005:
                notes.append(f"{name}: {aid} {tn} gained {(g - e) * 1000:.0f} mg since it was "
                             "weighed empty — something in it, or another rack's tube?")
    drift = {n: statistics.median(r) for n, r in ratios.items()}
    drifted: list[tuple[str, float]] = []        # files whose control tubes moved, mg
    shift = {n: statistics.median(r) for n, r in shifts.items()}
    mode = {True: "scale", False: ""}.get(study.drift_fix, study.drift_fix)  # a bool before
    moved = [(n, shift[n] * 1000, (r - 1) * 100) for n, r in drift.items()
             if abs(shift[n]) >= 0.0005]             # 0.5 mg, the balance's own repeatability
    drifted += [(n, mg) for n, mg, _ in moved if not mode]
    if moved:
        notes.append((f"{moved[0][0]}: the control tubes weigh {moved[0][1]:+.1f} mg "
                      f"({moved[0][2]:+.3f} %)" if len(moved) == 1 else
                      f"the control tubes weigh {min(x[1] for x in moved):+.1f} to "
                      f"{max(x[1] for x in moved):+.1f} mg in {len(moved)} files (" + ", ".join(
                          f"{_stem(n)[-15:]} {mg:+.1f} mg, {pct:+.3f} %" for n, mg, pct in moved)
                      + ")") + " against their empty weighing — " + (
                          "every tube corrected by them" if mode else
                          "balance drift? Results ▸ side panel ▸ control tubes"))
    if mode:
        for c in cells.values():
            c.fulls = [(n, (g / drift[n] if mode == "scale" else g - shift[n])
                        if k == "filled" and n in drift else g, k) for n, g, k in c.fulls]

    # ---- activities from file sources, every one corrected to `ref`
    covered: dict[str, set] = {}                 # count file -> the cells it counted
    spans: dict[str, tuple] = {}
    no_eff: dict[str, list[str]] = {}            # counter|window -> files with no efficiency
    for s in live:
        run = runs.get(s.path)
        if not run or s.kind in ("empty", "filled"):
            continue
        w = study.window_for(run)
        if not w:
            continue
        eff, own = study.eff_for(run, w)
        if eff is None:
            no_eff.setdefault(f"{run.counter}|{w}", []).append(Path(s.path).name)
        norm = run.normalized_to if own else None    # the counter's Bq at its reference
        mapping = s.mapping(run)
        for slot in run.slots:
            key = mapping.get(slot.key)
            if not key or not key[0] or w not in (slot.bq if own else slot.cpm):
                continue
            t_from = norm or slot.time
            if t_from is None:
                continue
            hl = study.hl_for(key[0])
            if eff is None:                      # CPM only, until an efficiency is typed
                cell(key).raw[Path(s.path).name] = (slot.counts.get(w), slot.cpm.get(w))
                continue
            bq = decay(slot.bq[w] if own else slot.cpm[w] / 60 / eff, t_from, ref, hl)
            stat = (slot.cpm if study.min_basis == "cpm" else slot.counts).get(w)
            cell(key).alts.append((Path(s.path).name, bq, slot.dead_time, stat, slot.time,
                                   rsd(slot, w)))
            cell(key).raw[Path(s.path).name] = (slot.counts.get(w), slot.cpm.get(w))
            covered.setdefault(Path(s.path).name, set()).add(key)
            if slot.time:
                t0, t1 = spans.get(Path(s.path).name, (slot.time, slot.time))
                spans[Path(s.path).name] = (min(t0, slot.time), max(t1, slot.time))
            cell(key).windows += [(Path(s.path).name, short, decay(slot.bq[x], t_from, ref, hl))
                                  for x, short in zip(run.windows, run.short_windows())
                                  if x in slot.bq and own]
            for x, short in zip(run.windows, run.short_windows()):
                e2, own2 = study.eff_for(run, x)
                if e2 and x in (slot.bq if own2 else slot.cpm):
                    cell(key).every.append((
                        Path(s.path).name, decay(slot.bq[x] if own2 else slot.cpm[x] / 60 / e2,
                                                 (run.normalized_to if own2 else None)
                                                 or slot.time, ref, hl),
                        slot.dead_time, (slot.cpm if study.min_basis == "cpm" else
                                         slot.counts).get(x), slot.time, short, slot.cpm.get(x),
                        x == w, slot.counts.get(x), rsd(slot, x)))

    for k, files in no_eff.items():
        counter, w = k.split("|", 1)
        ones = any(abs((runs[s.path].efficiency.get(w) or 0) - 1) < 1e-9 for s in live
                   if Path(s.path).name in files)
        notes.append(f"{counter} {w}: " + ("efficiency 1 in the file — none was set in the "
                                           "counter, its Bq are counts per second" if ones else
                                           "no counting efficiency in the file")
                     + f" ({len(files)} file(s)): activities in CPM only until it is typed — "
                     "Data sources ▸ efficiencies")

    # ---- counting rounds: the count files in time order, a new round whenever a vial comes
    # back — one round is one pass of the counter over the vials
    rounds: list[list[str]] = []
    seen: set = set()
    for name in sorted(covered, key=lambda n: spans.get(n, (_dt.datetime.max,))[0]):
        if not rounds or covered[name] & seen:
            rounds.append([])
            seen = set()
        rounds[-1].append(name)
        seen |= covered[name]
    round_of = {n: i for i, r in enumerate(rounds) for n in r}

    # ---- a blank's vial that reads background in one round and hot in another held something
    # else that time (260930: the tails, counted later in the bare-vial slot): not the blank's
    blanks_now = {t.name for t in study.tissues if t.role == "blank"}
    foreign: dict[str, list[str]] = {}
    for (aid, tn), c in cells.items():
        quiet = [x for x in c.alts if (x[3] or 0) < BLANK_COUNTS]
        hot = [x[0] for x in c.alts if (x[3] or 0) >= BLANK_COUNTS]
        if tn in blanks_now and quiet and hot:
            c.alts = [x for x in c.alts if x[0] not in hot]
            c.fulls = [x for x in c.fulls if x[0] not in hot]
            c.raw = {k: v for k, v in c.raw.items() if k not in hot}
            c.windows = [x for x in c.windows if x[0] not in hot]
            foreign.setdefault(tn, []).append(f"{aid} ({', '.join(_stem(n)[-15:] for n in hot)})")
    notes += [f"{tn} of {', '.join(xs)}: counted hot where it read background before — something "
              "else was in the vial that time (a tail counted later?): left out of the blank. To "
              "use it, place that vial on a tissue of its own (Data sources, unfold the file)"
              for tn, xs in foreign.items()]

    # ---- one activity per cell. A counting is valid with enough counts (or CPM) and a dead
    # time not too high (`is_valid`): one that is not is flagged, and used only when none of
    # the vial's is. The rule picks among those in the target range — more counts, a lower
    # dead time, a CPM not too high (a hot vial's first counting is out) —: the first, the
    # last, the one with the most counts, all combined (`combined`), one round's. None in
    # the target range: the rule among the valid ones; none valid: the nearest to valid —
    # the lowest dead time among those with enough counts, else the most counts.
    def valid(c):
        return [x for x in c.alts if is_valid(x, study)]

    def in_target(c, x):
        return ((x[2] or 1.0) <= study.dt_max and (x[3] is None or x[3] >= study.min_counts)
                and not (study.cpm_max and (c.raw.get(x[0], (0, 0))[1] or 0) > study.cpm_max))

    def rank(x):
        n = x[3] if x[3] is not None else math.inf
        return (0, x[2] or 1.0, -n) if n >= study.valid_counts else (1, -n, 0)

    def when(x):
        return x[4] or _dt.datetime.min

    def in_round(c, src):
        """The cell's counting from `src`, or from the round `src` is in."""
        return next((x for x in c.alts if x[0] == src), None) or next(
            (x for x in c.alts if src in round_of and round_of.get(x[0]) == round_of[src]), None)

    def in_window(c, src):
        """A counting in another window than the rule's: 'file@112-168'."""
        f, w = src.rsplit("@", 1)
        x = next((x for x in c.every if x[0] == f and x[5] == w), None)
        return x and (src, *x[1:5])

    def use(c, x, files):
        c.bq, c.bq_src, c.dead_time, c.counts = x[1], x[0], x[2], x[3]
        c.bq_used = files
    for c in cells.values():
        c.rule_w = {y[0]: y[5] for y in c.every if y[7]}
    rule_round = study.pick_count[6:] if study.pick_count.startswith("round:") else ""
    for c in cells.values():
        if c.alts:
            ok = valid(c)
            if study.count_agree:                    # out of the consensus: out
                ok = agreeing(ok, lambda a, b: counts_agree(a[1], a[5], b[1], b[5], study)) or ok
            ok = [x for x in ok if in_target(c, x)] or ok
            if study.pick_count == "all" and len(ok) > 1:
                use(c, combined(ok, study.combine), [x[0] for x in ok])
                continue
            if rule_round and (x := in_round(c, rule_round)):
                pick = x
            elif not ok:
                pick = min(c.alts, key=rank)         # none valid: the nearest to valid
            else:
                pick = (min(ok, key=when) if study.pick_count == "first" else
                        max(ok, key=when) if study.pick_count == "last" else
                        max(ok, key=lambda x: x[3] if x[3] is not None else math.inf))
            use(c, pick, [pick[0]])

    # ---- one mass per cell: filled minus empty tube, the filled tubes before a weigh+count
    # total (a recount's tubes may have lost their tissue), a mass typed by hand first
    picked = {(x[0], x[1], x[2]): x[3:] for x in study.chosen}   # what -> the files picked
    w_rank = {"filled": 2, "total": 1, "direct": 0}
    excess: dict[str, list[tuple]] = {}   # animal -> (tissue, day-of − later mg, later file)
    suggest: list[dict] = []              # what to offer the user: key, text, picks / fields
    lonely: dict[str, int] = {}
    light: dict[str, list[str]] = {}      # tissue -> where a tube weighed less than empty
    for key, c in cells.items():
        want = picked.get((*key, "empty"), [])
        es = [x for x in c.empties if x[0] in want] or c.empties[-1:]
        e = (" + ".join(x[0] for x in es), statistics.median(x[1] for x in es)) if es else None
        c.mass_empty = e[0] if e else ""
        usable = []
        for name, g, kind in c.fulls:
            if kind == "direct":         # typed by hand, or tared by the counter itself
                usable.append((name, g, (3 if study.pick_mass != "files" else -1)
                               if name == "manual" else 2.5))
            elif e is None:
                lonely[name] = lonely.get(name, 0) + 1
            elif g - e[1] < -0.005:                  # a bare vial, its tube left out of the run
                light.setdefault(key[1], []).append(f"{key[0]} ({_stem(name)[-15:]})")
            else:
                usable.append((name, g - e[1], w_rank[kind]))
        c.mass_alts = [(n, g) for n, g, _ in usable]
        # the first of the best-ranked: the day-of weighing, as the lab's sheets use. The tubes
        # are capped: a later weighing that differs is the day-of one being off (260930: cold
        # or wet tubes, +5-18 mg), or a tube moved — both flagged, the user picks. Several
        # weighings picked: their mean
        tubes = [u for u in usable if u[2] >= 1 and u[0] != "manual"]   # in file order
        want = picked.get((*key, "mass"), [])
        old = want[0] if len(want) == 1 and want[0] in ("mean", "median") else ""  # < 10.6
        how = old or study.mass_rule
        mine = [u for u in usable if u[0] in want]
        typed = next((u for u in usable if u[2] == 3), None)   # typed by hand, and it wins
        pool = (agreeing(tubes, lambda a, b: agree(a[1], b[1], study)) if study.mass_agree
                else []) or tubes                    # out of the consensus: out
        x = (mine[0] if len(mine) == 1 else
             (f"mean of {len(mine)}", statistics.fmean(u[1] for u in mine), 1) if mine else
             None) or (typed if not want else None) or (
            (f"{how} of {len(pool)}", _avg(how, [u[1] for u in pool]), 1)
            if how in ("mean", "median") and len(pool) > 1 else
            pool[-1] if how == "last" and pool else
            max(pool, key=lambda u: u[2]) if pool else
            max([u for u in usable if u[2] >= 0], key=lambda u: u[2], default=None))
        if x:
            c.mass_src, c.mass_g = x[0], x[1]
            c.mass_used = [u[0] for u in mine] if len(mine) > 1 else \
                [u[0] for u in pool] if x[0].startswith(("mean", "median")) else [x[0]]
            day = [u for u in usable if u[2] == 2]
            later = [u for u in usable if u[2] == 1]
            if day and later and 0 < day[0][1] < 0.2:
                excess.setdefault(key[0], []).append(
                    (key[1], (day[0][1] - later[0][1]) * 1000, later[0][0]))
    notes += [f"{name}: {n} filled tube(s) with no empty weight on the same cell — no mass"
              for name, n in lonely.items()]
    notes += [f"{tn} of {', '.join(xs)}: weighs less than its empty tube — no tube in that vial "
              "(a bare vial), that weighing not used" for tn, xs in light.items()]
    for aid, xs in excess.items():               # 260930: animals 1, 3, 5, 6, +8 to +18 mg
        heavy = sorted(((t, d, n) for t, d, n in xs if d > study.mass_tol_mg),
                       key=lambda x: -x[1])
        if len(heavy) >= 3 and statistics.median(d for _, d, _ in xs) > 3:
            mid = statistics.median(d for _, d, _ in heavy)
            suggest.append({"key": f"heavier:{aid}", "animal": aid, "mg": mid,
                            "picks": [[aid, t, "mass", n] for t, _, n in heavy],
                            "tissues": [f"{t} +{d:.0f} mg" for t, d, _ in heavy]})
            notes.append(f"{aid}: its small tissues weighed {mid:.0f}"
                         f" mg more on the day than when weighed again ("
                         + ", ".join(f"{t} +{d:.0f}" for t, d, _ in heavy[:6])
                         + ") — a like amount on each tube, the activity unchanged: something on "
                         "the outside that has since gone (condensation on cold tubes, water "
                         "from handling), or tubes weighed cold? Results ▸ select the cells ▸ "
                         "Mass ▸ the later weighing")

    # ---- an activity read on the dose calibrator wins over the counter
    for m in study.manual:
        t = study.tissue(m.tissue)
        if not m.animal or not m.tissue or m.mbq is None or (t and "activity" in t.closed):
            continue
        c = cell((m.animal, m.tissue))
        t = parse_time(m.time, study.day)
        inj = parse_time(a.inj_time, study.day) if (a := study.animal(m.animal)) else None
        if t is None:
            notes.append(f"{m.animal}/{m.tissue}: dose-calibrator time {m.time!r} not understood")
        else:
            if inj and t < inj:
                notes.append(f"{m.animal}/{m.tissue}: dose calibrator read at {t:%d %b %H:%M}, "
                             f"before the injection — another day? type the date too")
            c.calib = c.bq = decay(m.mbq * 1e6, t, ref, study.hl_for(m.animal))
            c.bq_src, c.dead_time, c.bq_used = "dose calibrator", None, ["dose calibrator"]

    # ---- countings picked by hand for some cells beat every rule: one, or several combined
    for (aid, tis, what), srcs in picked.items():
        c = cells.get((aid, tis))
        if not c or what != "count":
            continue
        if "dose calibrator" in srcs and c.calib is not None:
            c.bq, c.bq_src, c.dead_time, c.bq_used = c.calib, "dose calibrator", None, \
                ["dose calibrator"]
            continue
        xs = [x for s in srcs if (x := (in_window(c, s) if "@" in s else in_round(c, s)
                                         if s not in ("mean", "median") else None))]
        if srcs in (["mean"], ["median"]) and len(valid(c)) > 1:   # picked before 2026.10.6
            xs = valid(c)
        xs = list({x[0]: x for x in xs}.values())
        if xs:
            use(c, xs[0] if len(xs) == 1 else combined(xs, study.combine), [x[0] for x in xs])

    # ---- injected dose per animal
    injected, tails = {}, {}
    tail_names = [t.name for t in study.tissues if t.role == "tail"]
    for a in study.animals:
        notes += [f"{a.id}: {why} ({getattr(a, f)!r}) — another day? type the date too"
                  for f, why in time_problems(a, study.day).items()]
        hl = study.hl_for(a.id)
        tf, te = parse_time(a.full_time, study.day), parse_time(a.empty_time, study.day)
        if a.full_mbq is None or tf is None:
            continue
        idb = decay(a.full_mbq * 1e6, tf, ref, hl)
        if a.empty_mbq is not None and te is not None:
            idb -= decay(a.empty_mbq * 1e6, te, ref, hl)
        for lo in a.losses:
            mbq, tl = _float(lo.get("mbq")), parse_time(lo.get("time", ""), study.day)
            if mbq is None:
                continue
            if tl is None:
                notes.append(f"{a.id}: other loss time {lo.get('time')!r} not understood")
            else:
                idb -= decay(mbq * 1e6, tl, ref, hl)
        tail = sum(cells[(a.id, n)].bq for n in tail_names
                   if (a.id, n) in cells and cells[(a.id, n)].bq)
        empty_vial = abs(tail) < 1e-4 * idb and any(                # counter background
            cells[(a.id, n)].bq is not None for n in tail_names if (a.id, n) in cells)
        if a.tail_mbq is not None:
            tt = parse_time(a.tail_time, study.day)
            if tt is None:
                notes.append(f"{a.id}: tail time {a.tail_time!r} not understood")
            else:
                vial, tail = tail, decay(a.tail_mbq * 1e6, tt, ref, hl)
                if not empty_vial and vial and abs(vial - tail) > 0.2 * max(vial, tail):
                    notes.append(f"{a.id}: tail on the card {tail / 1e6:.3g} MBq but the tail "
                                 f"vial reads {vial / 1e6:.3g} MBq (at the reference) — the "
                                 f"card is used")
        elif empty_vial:
            notes.append(f"{a.id}: the tail vial reads {tail / 1e6:.1g} MBq, background — "
                         f"read on the dose calibrator? type it on the card")
        tails[a.id] = tail
        if study.subtract_tail:
            idb -= tail
        injected[a.id] = idb

    # ---- flags
    def namer(c):
        """A row of the side panel by what it is: 'round 2', 'count+weight r3', 'weight'."""
        kinds = {n: k for n, _, k in c.fulls}

        def name(key):
            i = round_of.get(key[1])
            r = f"round {i + 1}" if i is not None else _stem(key[1])[-15:]
            if key[0] == "count":
                return r + ("" if c.rule_w.get(key[1]) == key[2] else f" · {key[2]}")
            if key[0] == "mass":
                return "weight" if kinds.get(key[1]) == "filled" else \
                    f"count+weight r{i + 1}" if i is not None else _stem(key[1])[-15:]
            return "tare " + _stem(key[1])[-6:]
        return name
    for (aid, tname), c in cells.items():
        t = study.tissue(tname)
        if c.mass_g is not None and c.mass_g <= 0 and (not t or t.role != "blank"):
            c.flags.append("mass <= 0")
        if c.bq is not None and c.bq <= 0 and (not t or t.role != "blank"):
            c.flags.append("at background")
        if t and t.role == "blank" and (c.counts or 0) >= BLANK_COUNTS:
            c.flags.append("activity in a blank")
        if c.dead_time and c.dead_time > study.valid_dt:
            c.flags.append(f"dead time {c.dead_time:.2f}")
        if c.bq_src != "dose calibrator" and c.counts is not None and \
                c.counts < study.valid_counts and (not t or t.role != "blank"):
            c.flags.append(f"low counts: {c.counts:,.0f} {study.min_basis}")
        c.pairs = pairs_of(c, study, namer(c))
        used = [("count", *(u.rsplit("@", 1) if "@" in u else (u, c.rule_w.get(u, ""))))
                for u in c.bq_used]
        if gaps := [g for u in used for g in c.pairs.get(u, [])]:
            c.flags.append(f"counts differ: {gaps[0]}" + (f" (+{len(gaps) - 1})"
                                                          if len(gaps) > 1 else ""))
        if gaps := [g for u in c.mass_used for g in c.pairs.get(("mass", u), [])]:
            c.flags.append(f"weighings differ: {gaps[0]}" + (f" (+{len(gaps) - 1})"
                                                             if len(gaps) > 1 else ""))
        if c.bq is not None and c.mass_g is None and (not t or t.role in ("tissue", "tail")):
            c.flags.append("no mass")

    if drifted:
        suggest.append({"key": "drift:" + ",".join(n for n, _ in drifted), "drift": drifted})
    refs = {a.id: t for a in study.animals if study.ref_rule == "injection"
            and (t := parse_time(a.inj_time, study.day))}
    res = Result(ref=ref, cells=cells, injected_bq=injected, tail_bq=tails, notes=notes,
                  suggest=suggest,
                  refs=refs, hls={a.id: study.hl_for(a.id) for a in study.animals},
                  rounds=rounds, spans=spans,
                  files={Path(s.path).name: (s.kind, runs[s.path].started) for s in live
                         if s.path in runs}, no_eff=list(no_eff))
    range_flags(study, res, study.ranges)
    return res


def same_tissue(a: str, b: str) -> bool:
    """'Thyroids', 'thyroid', 'Thyr': one tissue named two ways (case, a plural, a short
    form of at least four letters)."""
    a, b = (re.sub(r"s$", "", x.strip().lower()) for x in (a, b))
    return bool(a and b) and (a == b or min(len(a), len(b)) >= 4 and
                              (a.startswith(b) or b.startswith(a)))


def range_flags(study: Study, res: Result, ranges) -> None:
    """Flag the tissues outside what is expected of them: ranges = [[tissue, mg min, mg max,
    %IA/g min, %IA/g max]] (Options; a tissue named a little differently still matches,
    see `same_tissue`), 0 = no bound."""
    for (aid, tn), c in res.cells.items():
        t = study.tissue(tn)
        rule = next((r for r in ranges if same_tissue(r[0], tn)), None)
        if not rule or not t or t.role != "tissue":
            continue
        for v, lo, hi, what in ((None if c.mass_g is None else c.mass_g * 1000, rule[1],
                                 rule[2], "mg"),
                                (res.value(study, aid, tn, "pid_g"), rule[3], rule[4], "%IA/g")):
            if v is not None and (lo and v < lo or hi and v > hi):
                c.flags.append(f"out of range: {v:.3g} {what} ({lo:g}–{hi:g})")


def _avg(how: str, vals: list[float]) -> float:
    return statistics.median(vals) if how == "median" else statistics.fmean(vals)


def agree(a: float, b: float, study: Study) -> bool:
    """Two weighings (g) agree: within the study's mg, or its % of the lighter."""
    d = abs(a - b) * 1000
    return d <= study.mass_tol_mg or d <= study.mass_tol_pct / 100 * 1000 * min(abs(a), abs(b))


def is_valid(x: tuple, study: Study) -> bool:
    """A counting (src, Bq, dead time, counts or CPM, …) enough counted and not too busy to
    be taken: else flagged, and used only when no other counting of the vial is valid."""
    return (x[2] or 1.0) <= study.valid_dt and (x[3] is None or x[3] >= study.valid_counts)


def rsd(slot: hidex.Slot, w: str) -> float | None:
    """A counting's relative σ from counting alone: √(counts as counted) over the counts the
    value stands on — background off (the Hidex takes it off) and dead time undone: CPM ×
    time / 60 / dead time. 260930 6/Gall bladder round 3: 219 counted, 57 of them the
    tissue's — ±26 %, not the ±7 % of 219. No counted time in the file: √counts / counts."""
    n, cpm = slot.counts.get(w), slot.cpm.get(w)
    if not n or n <= 0:
        return None
    net = cpm * slot.secs / 60 / (slot.dead_time or 1.0) if slot.secs and cpm is not None else n
    return math.sqrt(n) / net if net > 0 else math.inf


def count_gap(a: float, ra, b: float, rb) -> tuple[float, float]:
    """Two countings of a vial at one instant (Bq), b against a: in %, and in σ of their
    counting statistics — ra, rb each one's relative σ (`rsd`); inf when not known."""
    pct = (b - a) / a * 100
    z = abs(b - a) / math.hypot(a * ra, b * rb) if ra and rb else math.inf
    return pct, z


def counts_agree(a: float, ra, b: float, rb, study: Study) -> bool:
    pct, z = count_gap(a, ra, b, rb)
    return abs(pct) <= study.count_tol_pct or z <= study.count_tol_sigma


def agreeing(items: list[tuple], same) -> list[tuple]:
    """The measures of one vial (its weighings, its countings) that agree with each other
    — `same(a, b)` —, when more than half do: the largest such set, the earliest first
    among equals (3 weighings, 2 agree: the other is left out; 2 that disagree: none).
    ponytail: every subset — a vial is measured a handful of times, not dozens."""
    for size in range(len(items), 1, -1):
        if size * 2 <= len(items):
            break
        for grp in itertools.combinations(items, size):
            if all(same(a, b) for a, b in itertools.combinations(grp, 2)):
                return list(grp)
    return []


def pairs_of(c: Cell, study: Study, name=None) -> dict[tuple, list[str]]:
    """Every two measures of one vial that should agree and do not: valid countings in the
    same window that do not agree within the study's % or σ of their counting statistics
    (`counts_agree`); weighings of the filled tube (the tissue mass they give) and tares that
    do not `agree`. When more than half agree (`agreeing`, the consensus), each one out of it
    is told so, and against which; one in it is not told of those out of it — a thin late
    counting off by chance does not flag the good ones. No consensus: both are told (13.png:
    a day-of weighing 2 mg over two later ones that agree is the one out, flagged; they are
    not). {row key: ["out of the consensus …", "+8.1 % against round 2 (4.2 σ)", …]}; `name`
    says what a row key is ("round 2", "count+weight r3"), else its file."""
    out: dict[tuple, list[str]] = {}
    name = name or (lambda key: _stem(key[1])[-15:])

    def told(groups):
        """groups: [(items, an item's row key, same(a, b), the text of a against b)]."""
        for items, key, same, gap in groups:
            cons = {key(x) for x in agreeing(items, same)}
            for x in items:
                if cons and key(x) not in cons:
                    out.setdefault(key(x), []).append(
                        f"out of the consensus — {len(cons)} of {len(items)} agree, not this one")
            for i, a in enumerate(items):
                for b in items[i + 1:]:
                    if not same(a, b):
                        for x, y in ((a, b), (b, a)):
                            if key(x) not in cons:
                                out.setdefault(key(x), []).append(gap(x, y))
    xs = [y for y in c.every if is_valid(y, study) and y[1] > 0]

    def cgap(a, b):
        pct, z = count_gap(b[1], b[9], a[1], a[9])   # a against b
        return f"{pct:+.1f} % against {name(('count', b[0], b[5]))}" + (
            f" ({z:.1f} σ)" if z != math.inf else "")
    groups = [([y for y in xs if y[5] == w], lambda y: ("count", y[0], y[5]),
               lambda a, b: a[0] == b[0] or counts_agree(a[1], a[9], b[1], b[9], study), cgap)
              for w in dict.fromkeys(y[5] for y in xs)]
    e = statistics.median(g for n, g in c.empties if n in c.mass_empty.split(" + ")) \
        if c.mass_empty else None
    ws = [(n, g if k == "direct" else g - e) for n, g, k in c.fulls
          if n != "manual" and (k == "direct" or e is not None)]
    for what, items in (("mass", ws), ("tare", c.empties)):
        groups.append(([x for x in items if x[1] > -0.005], lambda x, w=what: (w, x[0]),
                       lambda a, b, w=what: agree(a[1], b[1], study) if w == "mass" else
                       abs(a[1] - b[1]) * 1000 <= study.mass_tol_mg,   # a tube: mg, not %
                       lambda a, b, w=what: f"{(a[1] - b[1]) * 1000:+.1f} mg against "
                                            f"{name((w, b[0]))}"))
    told(groups)
    return out


def _stem(name: str) -> str:
    """A file as the notes and the side panel name it: 'Tc-99m-002-20261001-090037'."""
    return re.sub(r"(-AutoExport)?\.(xlsx|csv)$", "", str(name))


def combined(xs: list[tuple], how: str = "mean") -> tuple:
    """Several countings of one vial as one, their activities all decay-corrected to the same
    instant: their mean, or "weighted" by their counts (as counted — the inverse of each one's
    relative Poisson variance; decay correction scales a value, not how sure it is). With
    70,968 and 3,406 counts the mean gives each half, weighted 95 % and 5 %."""
    w = [x[3] or 0 for x in xs] if how == "weighted" else []
    w = w if w and all(w) else [1] * len(xs)
    return (f"{'weighted' if how == 'weighted' else 'mean'} of {len(xs)}",
            sum(x[1] * k for x, k in zip(xs, w)) / sum(w),
            max((x[2] for x in xs if x[2]), default=None), sum(x[3] or 0 for x in xs), None)


def _float(text) -> float | None:
    try:
        return float(str(text).strip().replace(",", "."))
    except ValueError:
        return None


def at_injection(study: Study, res: Result, aid: str, bq: float | None) -> float | None:
    """An activity at the reference time, taken back to the animal's injection time."""
    a = study.animal(aid)
    t = parse_time(a.inj_time, study.day) if a else None
    return None if bq is None or t is None else decay(bq, res.ref, t, study.hl_for(aid))


def at_imaging(study: Study, res: Result, aid: str) -> list[tuple[str, _dt.datetime, float]]:
    """The activity in the animal at the start of each SPECT / PET session: (modality, start,
    Bq) — the injected activity (tail off, as set) decayed to the session's start."""
    a, idb = study.animal(aid), res.injected_bq.get(aid)
    if not a or idb is None:
        return []
    out = []
    for ev in a.events:
        mod = str(ev.get("modality") or ev.get("what") or "").strip()
        t = event_time(a, ev.get("start") or ev.get("when", ""), study.day)
        if ev.get("kind") != "imaging" or not t or mod.lower() in ("ct", "mri", "mr"):
            continue                             # CT, MRI alone: no activity to speak of
        name = ("SPECT" if re.search(r"(?i)spect|scinti", mod) else
                "PET" if re.search(r"(?i)\bpet", mod) else "imaging")
        out.append((name, t, decay(idb, res.ref, t, study.hl_for(aid))))
    return out


def imaging_label(study: Study, res: Result) -> str:
    """'SPECT', 'PET', 'SPECT / PET' or 'imaging': what the sessions of the study were."""
    names = sorted({m for a in study.animals for m, _, _ in at_imaging(study, res, a.id)})
    return " / ".join(names) or "SPECT / PET"


def tail_pct(study: Study, res: Result, aid: str) -> float | None:
    """The injection site as a share of what left the syringe, in %IA."""
    tail, idb = res.tail_bq.get(aid), res.injected_bq.get(aid)
    if tail is None or idb is None:
        return None
    dose = idb + (tail if study.subtract_tail else 0)
    return tail / dose * 100.0 if dose > 0 else None


def recovery_pct(study: Study, res: Result, aid: str) -> float | None:
    """The injected activity found in all the collected tissues together, in %IA — a blunt
    sanity number: well under 100 is normal (carcass, excreta), far over is a mis-mapped run."""
    idb = res.injected_bq.get(aid)
    if not idb:
        return None
    tot = sum(c.bq for (a, n), c in res.cells.items()
              if a == aid and c.bq and (t := study.tissue(n)) and t.role == "tissue")
    return tot / idb * 100.0


# ------------------------------------------------------------- guessing the sources
def strip_summary(pairs: list) -> str:
    """A file's vials in one line, in rack order: '107: Adrenal, BAT … WAT · 108: …' or
    'tumor: 107, 108, 109 · muscle: …'; a None pair is a vial not used."""
    def names(xs):
        return ", ".join(xs) if len(xs) <= 4 else f"{xs[0]}, {xs[1]} … {xs[-1]}"
    out, i = [], 0
    while i < len(pairs):
        p = pairs[i]

        def end(k):
            return next((j for j in range(i, len(pairs)) if (pairs[j] or (None, None))[k]
                         != (p or (None, None))[k]), len(pairs))
        if not p:
            j = end(0)
            out.append(f"{j - i} not used")
        elif (j := end(0)) >= end(1):
            out.append(f"{p[0]}: {names([x[1] for x in pairs[i:j]])}")
        else:
            j = end(1)
            out.append(f"{p[1]}: {names([x[0] for x in pairs[i:j]])}")
        i = j
    return " · ".join(out)


def names_summary(used: list[str], every: list[str]) -> str:
    """The names a file uses, against the list they come from: 'all', 'all but BAT, WAT',
    else the names themselves, in the list's order."""
    used = [x for x in every if x in used] + [x for x in dict.fromkeys(used) if x not in every]
    missing = [x for x in every if x not in used]
    if not used:
        return "none"
    if not missing:
        return "all"
    if len(missing) <= 3 and len(used) > len(missing):
        return "all but " + ", ".join(missing)
    return ", ".join(used)


def propose(vials: list, c: int, animals: list[str], tissues: list[str], order: str) -> dict:
    """Vial `c` just placed by hand: how the vials after it go on, {index: [animal, tissue]}
    for those that would change. The order follows the vial before it — same animal: animal
    by animal; same tissue: tissue by tissue — else `order`."""
    here = vials[c] or ["", ""]
    prev = (vials[c - 1] or ["", ""]) if c else ["", ""]
    if not (len(here) > 1 and here[0] in animals and here[1] in tissues):
        return {}
    if len(prev) > 1 and prev[0] == here[0] and prev[1] != here[1]:
        order = ANIMAL_MAJOR
    elif len(prev) > 1 and prev[1] == here[1] and prev[0] != here[0]:
        order = TISSUE_MAJOR
    pairs = ([[a, t] for t in tissues for a in animals] if order == TISSUE_MAJOR
             else [[a, t] for a in animals for t in tissues])
    nxt = pairs[pairs.index(list(here[:2])) + 1:]
    return {k: p for k, p in zip(range(c + 1, len(vials)), nxt)
            if list(vials[k] or []) != p}


def _place(s: Source, run: hidex.Run, want: dict[str, tuple[str, str]], names: list[str],
           preset: str = ANIMAL_MAJOR):
    """Write a slot -> (animal, tissue) placement as a preset plus the exceptions, so the
    table still reads as animals x tissues."""
    s.preset, s.slotmap, s.tissues = preset, {}, list(names)
    s.animals = list(dict.fromkeys(a for a, _ in want.values()))
    base = s.mapping(run)
    s.slotmap = {k: list(v) for k, v in want.items() if base.get(k) != v}
    s.slotmap.update({k: [] for k in base if k not in want})


def _match_racks(blocks) -> list[tuple[int, int]]:
    """(empty, filled) pairs among weighed racks. Filled tubes weigh at least what they did
    empty (5 mg of balance slack) and most gained something; a true pair keeps its smallest
    gain near zero (a blank tube, a tiny sample) while a wrong pair goes negative on some
    tube, so the smallest non-negative gains are taken first. A run whose every rack matches
    the same-place rack of one other run goes first: animals are weighed whole far more often
    than their racks are shuffled. One tube of a rack may disagree (a tube taken out between
    the two weighings — 260930: a bare vial weighed with a tube in it the first time): a
    match with such a tube ranks after the clean ones.
    ponytail: greedy, not an assignment solver — fine for a few dozen racks."""
    runs: dict[int, list[int]] = {}
    for i, (r, _) in enumerate(blocks):
        runs.setdefault(id(r), []).append(i)
    cands = []
    for i, (ra, a) in enumerate(blocks):
        for j, (rb, b) in enumerate(blocks):
            if len(a) != len(b) or not (ra.started and rb.started and ra.started < rb.started):
                continue
            d = [y.tare_g - x.tare_g for x, y in zip(a, b)
                 if x.tare_g is not None and y.tare_g is not None]
            odd = [g for g in d if g <= -0.005]
            rest = [g for g in d if g > -0.005]
            if len(d) == len(a) and len(odd) <= (len(d) >= 4) and rest \
                    and statistics.median(d) > 0.010:
                cands.append((len(odd), min(rest) < 0, abs(min(rest)), i, j))
    ok = {(i, j) for odd, *_, i, j in cands if not odd}   # a whole run: clean racks only

    def whole(i, j):
        a, b = runs[id(blocks[i][0])], runs[id(blocks[j][0])]
        return len(a) == len(b) and all(p in ok for p in zip(a, b))
    pairs, used, side = [], set(), {}     # side: run -> "empty" | "filled", never both
    for _, w, *_, i, j in sorted((c[0], not whole(c[-2], c[-1]), *c) for c in cands):
        ri, rj = id(blocks[i][0]), id(blocks[j][0])
        # an empty run matched whole again is the same tubes weighed a second time
        if (i in used and w) or j in used or side.get(ri, "empty") != "empty" \
                or side.get(rj, "filled") != "filled":
            continue
        used |= {i, j}
        side[ri], side[rj] = "empty", "filled"
        pairs.append((i, j))
    return pairs


def _recount_of(study: Study, run: hidex.Run, earlier: list[hidex.Run]) -> hidex.Run | None:
    """The earlier run these same vials were counted in: once decayed to the same instant,
    a recount agrees vial by vial within a few % (dead time, statistics) while two animals
    of one group differ by tens of %."""
    w = study.window_for(run)
    hl = run.half_life_s.get(w) or HALF_LIFE_S["99mTc"]

    def level(r, s, x, at):                      # Bq (Hidex) or CPM (Wizard2) at one instant
        if r.normalized_to:
            return decay(s.bq[x], r.normalized_to, at, hl) if s.bq.get(x, 0) > 0 else None
        return decay(s.cpm[x], s.time, at, hl) if s.cpm.get(x, 0) > 0 and s.time else None
    for e in earlier:
        we = study.window_for(e)
        if (len(e.slots) != len(run.slots) or not w or not we or not e.started
                or (e.normalized_to is None) != (run.normalized_to is None)):
            continue
        pairs = [(level(e, x, we, e.started), level(run, y, w, e.started))
                 for x, y in zip(e.slots, run.slots)
                 if min(x.counts.get(we, 0), y.counts.get(w, 0)) >= RECOUNT_COUNTS]
        lr = [abs(math.log(b / a)) for a, b in pairs if a and b]
        if len(lr) >= 3 and statistics.median(lr) < 0.1:
            return e
    return None


def auto_assign(study: Study, runs: dict[str, hidex.Run], why: dict | None = None) -> list[str]:
    """Guess what every `auto` source is and where its vials go; returns notes for the user,
    and fills `why` with each guessed file's reason (uid -> text) for the log.

    Weighings are matched rack by rack on tube weight: the earlier side of a match is the
    empty tubes, dealt to the animals in time order; a filled rack takes its places from
    the empty rack it matched, so a rack weighed out of turn still lands on its animal.
    Counts follow the animals in time order, except a run repeating an earlier run's
    activities, which is a recount of the same vials. Each run starts a new animal.

    Tissues counted in runs of their own (a `batch`: collected and counted on the spot) have
    their own runs: a run goes to the batch whose tissue count fits its vial count, and each
    batch deals its animals on its own, in its own vial order.
    """
    ids = [a.id for a in study.animals if a.id]
    batches = {b: study.batch_tissues(b) for b in study.batch_names()}
    batches = {b: n for b, n in batches.items() if n}
    srcs = sorted((s for s in study.sources
                   if s.path in runs and s.kind != "ignored"),
                  key=lambda s: runs[s.path].started or _dt.datetime.min)
    if not ids or not batches or not any(s.auto for s in srcs):
        return []
    notes: list[str] = []
    why = {} if why is None else why
    for s in srcs:
        if s.auto:                  # the batch whose tissues fill the run's vials best
            nv = len(runs[s.path].slots)
            s.batch = min(batches, key=lambda b: (nv % len(batches[b]) != 0,
                                                  nv % len(batches[b]), -len(batches[b])))

    def deal(todo, kinds):
        for b, names in batches.items():
            mine = [s for s in todo if s.batch == b]
            if not mine:
                continue
            nt, tm = len(names), study.batch_order(b) == TISSUE_MAJOR
            taken = {a for s in srcs if not s.auto and s.kind in kinds and s.batch == b
                     for a in s.animals}
            free, k = [a for a in ids if a not in taken], 0
            for s in mine:
                run = runs[s.path]
                n_an = -(-len(run.slots) // nt)       # a run starts on a fresh animal
                who = free[k:k + n_an]
                k += n_an
                cells = ([(a, t) for t in names for a in who] if tm else
                         [(a, t) for a in who for t in names])
                if tm and len(who) < n_an:            # short of animals: keep the grid shape
                    cells = [(a, t) for t in names for a in who + [None] * (n_an - len(who))]
                want = {sl.key: c for sl, c in zip(run.slots, cells) if c[0]}
                why[s.uid] = (f"{KIND_LABEL[s.kind]} file {mine.index(s) + 1} of {len(mine)} "
                              f"in time order → {', '.join(who) or 'no animal left'}")
                if len(want) < len(run.slots) and s is mine[-1]:
                    notes.append(f"{Path(s.path).name}: {len(run.slots) - len(want)} vial(s) "
                                 f"left over — more vials than animals x tissues")
                _place(s, run, want, names, TISSUE_MAJOR if tm else ANIMAL_MAJOR)

    # ---- weighings: empty or filled, by matching racks on tube weight
    tares = [s for s in srcs if runs[s.path].kind == "tare"]
    blocks, owner = [], []
    for s in tares:
        racks: dict[int, list[hidex.Slot]] = {}
        for sl in runs[s.path].slots:
            racks.setdefault(sl.rack, []).append(sl)
        for r, sls in racks.items():
            blocks.append((runs[s.path], sls))
            owner.append((s, r))
    pairs = _match_racks(blocks)
    for s in tares:
        if s.auto:
            if any(owner[j][0] is s for _, j in pairs):
                s.kind = "filled"
            elif any(owner[i][0] is s for i, _ in pairs):
                s.kind = "empty"
            elif pairs and not any(o.kind == "weigh_count" and len(runs[o.path].slots)
                                   == len(runs[s.path].slots) for o in srcs):
                # ponytail: a same-size count + weight run is taken for its filled side
                notes.append(f"{Path(s.path).name}: no rack matches another weighing by tube "
                             f"weight — check its kind")
    deal([s for s in tares if s.auto and s.kind == "empty"], ("empty",))
    matched: dict[str, list[str]] = {}           # weighing uid -> the files its racks matched
    for i, j in pairs:
        for s, o in ((owner[i][0], owner[j][0]), (owner[j][0], owner[i][0])):
            matched.setdefault(s.uid, [])
            if Path(o.path).name not in matched[s.uid]:
                matched[s.uid].append(Path(o.path).name)
    for s in tares:
        if s.auto and s.uid in matched:
            m = ", ".join(matched[s.uid])
            why[s.uid] = (f"tube weights: lighter than {m}, rack by rack; " + why.get(s.uid, "")
                          if s.kind == "empty" else
                          f"tube weights: heavier than the empty tubes of {m}, rack by rack — "
                          "each rack takes their places")

    for sf in tares:
        if not sf.auto or sf.kind != "filled":
            continue
        want, lead = {}, None
        for i, j in sorted((p for p in pairs if owner[p[1]][0] is sf), key=lambda p: owner[p[1]][1]):
            (se, re_), rf = owner[i], owner[j][1]
            if se.kind != "empty":
                continue
            emap = se.mapping(runs[se.path])
            got = [emap.get(e.key) for e in blocks[i][1]]
            want.update({f.key: g for f, g in zip(blocks[j][1], got) if g})
            lead = lead or se
            sf.batch = se.batch
            if rf != re_ or se is not lead:           # not where the running order puts it
                who = ", ".join(dict.fromkeys(g[0] for g in got if g))
                notes.append(f"rack slip: {Path(sf.path).name} rack {rf} holds the tubes of "
                             f"{who or '?'} ({Path(se.path).name} rack {re_}) — matched by "
                             f"tube weight and placed there")
        _place(sf, runs[sf.path], want, batches.get(sf.batch, []))
        if len(want) < len(runs[sf.path].slots):
            notes.append(f"{Path(sf.path).name}: {len(runs[sf.path].slots) - len(want)} "
                         f"filled tube(s) matched no empty rack")

    for i, j in pairs:                           # the tube a match let through
        for x, y in zip(blocks[i][1], blocks[j][1]):
            if x.tare_g is not None and y.tare_g is not None and y.tare_g - x.tare_g <= -0.005:
                notes.append(f"{Path(owner[j][0].path).name} vial {y.key}: {x.tare_g - y.tare_g:.3f}"
                             f" g lighter than its empty tube ({Path(owner[i][0].path).name}) — a "
                             f"tube taken out? the rack is matched on its other tubes")

    # ---- counts: in animal order, a recount goes where its first count went
    fresh, seen = [], []
    for s in (s for s in srcs if s.kind in ("count", "weigh_count")):
        run = runs[s.path]
        first = _recount_of(study, run, [runs[o.path] for o in seen])
        if s.auto and first:
            orig = next(o for o in seen if runs[o.path] is first)
            s.batch = orig.batch
            _place(s, run, orig.mapping(first), batches.get(s.batch, []), orig.preset)
            why[s.uid] = (f"a recount of {first.name}: the same activities vial by vial, "
                          "once decayed to the same time")
        elif s.auto:
            fresh.append(s)
            deal(fresh, ("count", "weigh_count"))
        seen.append(s)
    return notes


def unplaced(study: Study, runs: dict[str, hidex.Run]) -> list[str]:
    """The check behind the guess: every vial of every file in use lands on an animal and a
    tissue — a file with vials going nowhere, or a tube weighed empty whose filled weighing
    went nowhere (and the other way), is named."""
    out, placed = [], {"empty": set(), "filled": set()}
    for s in study.sources:
        run = runs.get(s.path)
        if s.kind == "ignored" or not run:
            continue
        m = s.mapping(run)
        lost = [sl.key for sl in run.slots if sl.key not in m]
        if lost:
            out.append(f"{Path(s.path).name}: {len(lost)} of {len(run.slots)} vials placed "
                       f"nowhere ({', '.join(lost[:8])}{' …' if len(lost) > 8 else ''}) — "
                       f"unfold it in Data sources")
        if s.kind in placed:
            placed[s.kind] |= set(m.values())
    if placed["empty"] and placed["filled"]:
        for a in study.animals:
            e = {t for x, t in placed["empty"] if x == a.id}
            f = {t for x, t in placed["filled"] if x == a.id}
            if e and f and e != f:
                miss = sorted(e ^ f)
                out.append(f"{a.label}: {', '.join(miss[:6])}{' …' if len(miss) > 6 else ''} "
                           f"weighed {'empty only' if miss[0] in e else 'filled only'} — "
                           f"no mass for {'it' if len(miss) == 1 else 'them'}")
    return out


def carry_fixes(study: Study, runs: dict[str, hidex.Run]) -> list[tuple[str, str, str, dict]]:
    """An animal's tubes put in order by hand in one file (a tissue skipped when they were
    filled: 260930, animal 2's gall bladder went in tube 9, not 3) are the same tubes, in the
    same order, in its other files. For every other file holding that animal's tissues in
    another order — and not placed by hand for it — (file uid, animal, the fixed file's
    name, {slot: [animal, tissue]}) to put them the same way."""
    live = [s for s in study.sources if s.kind != "ignored" and s.path in runs]

    def seq(s, a):
        m = s.mapping(runs[s.path])
        return [(sl.key, m[sl.key][1]) for sl in runs[s.path].slots
                if sl.key in m and m[sl.key][0] == a]

    def by_hand(s, a):
        return not s.auto and any(len(v) > 1 and v[0] == a for v in s.slotmap.values())
    out, seen = [], set()
    for s in live:
        for a in dict.fromkeys(v[0] for v in s.slotmap.values() if len(v) > 1 and v[0]):
            if not by_hand(s, a):
                continue
            want = [t for _, t in seq(s, a)]
            for o in live:
                got = seq(o, a) if o is not s and not by_hand(o, a) else []
                if (o.uid, a) in seen or len(got) != len(want) \
                        or [t for _, t in got] == want or sorted(t for _, t in got) != sorted(want):
                    continue
                seen.add((o.uid, a))
                out.append((o.uid, a, Path(s.path).name,
                            {k: [a, t] for (k, old), t in zip(got, want) if old != t}))
    return out


# ---------------------------------------------------------------- import helpers
def _grids(path) -> list[list[list]]:
    """Every sheet of a .xlsx (or the one of a .csv) as rows of cell values."""
    p = Path(path)
    if p.suffix.lower() in (".csv", ".txt"):
        text = p.read_text(encoding="utf-8-sig", errors="replace")
        rows = list(csv.reader(text.splitlines(), csv.Sniffer().sniff(text[:2000], ";,\t")))
        num = re.compile(r"-?\d+(?:[.,]\d+)?")
        return [[[float(v.replace(",", ".")) if num.fullmatch(v.strip()) else v.strip() or None
                  for v in r] for r in rows]]
    import openpyxl
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    return [[list(r) for r in ws.iter_rows(values_only=True)] for ws in wb.worksheets]


def _mass(nums: list[float]) -> float | None:
    """A tissue's mass from the numbers of one animal's block on its row: the difference a
    sheet works out (empty, full, full − empty), a mass weighed as it is, or full − empty."""
    for i, a in enumerate(nums):
        for j in range(i + 1, len(nums)):
            for k in range(j + 1, len(nums)):
                if abs(nums[j] - a - nums[k]) < 1e-6:
                    return nums[k]
    if len(nums) == 1 or nums and max(nums) - min(nums) < 1e-9:
        return nums[0]
    if len(nums) == 2 and min(nums) > 1:                    # two tube weighings
        return nums[1] - nums[0]
    return None


def read_weights(path, animals: list["Animal"], tissues: list[str]):
    """A sheet of masses made by hand: animals across (a row holding their IDs or aliases),
    tissues down (the study's names, left of each animal's block), in g. Returns
    ([(animal id, tissue, g)], rows not understood, labels not in the tissue list) — or
    None when it does not look like one (no animal and no tissue of the study on it)."""
    names = {}
    for a in animals:
        for n in [a.id, *a.aliases]:
            if n:
                names[str(n).strip().lower()] = a.id
    tis = {t.lower(): t for t in tissues}

    def text(v):
        return (f"{v:g}" if isinstance(v, float) else str(v)).strip() if v is not None else ""
    got, unsure, unknown = {}, [], []
    for g in _grids(path):
        heads = []
        for r, row in enumerate(g[:8]):
            hit = [(c, names[text(v).lower()]) for c, v in enumerate(row)
                   if text(v).lower() in names]
            if len(hit) > len(heads[1]) if heads else hit:
                heads = (r, hit)
        if not heads:
            continue
        hr, cols = heads
        marks = [c for c, v in enumerate(g[hr]) if text(v)] + [max(len(x) for x in g)]
        for row in g[hr + 1:]:
            for c0, aid in cols:                 # a block ends where the next heading is
                c1 = next(c for c in marks if c > c0)
                label = next((text(row[c]) for c in range(min(c0, len(row) - 1), -1, -1)
                              if isinstance(row[c], str) and text(row[c])), "")
                if not label:
                    continue
                if label.lower() not in tis:
                    unknown.append(label)
                    continue
                nums = [v for v in row[c0:c1] if isinstance(v, (int, float))
                        and not isinstance(v, bool)]
                if not nums:
                    continue
                m = _mass(nums)
                if m is None:
                    unsure.append(f"{aid} {label}: {', '.join(f'{v:g}' for v in nums)}")
                else:
                    got[(aid, tis[label.lower()])] = round(m, 6)
    if not got and not unsure:
        return None
    return ([(a, t, m) for (a, t), m in got.items()], unsure,
            list(dict.fromkeys(unknown)))


def read_column(path) -> list[str]:
    """First non-empty column of a one-column xlsx/csv — a tissue list."""
    p = Path(path)
    if p.suffix.lower() in (".csv", ".txt", ".tsv"):
        sep = "\t" if p.suffix.lower() == ".tsv" else ","
        rows = [r.split(sep)[0].strip() for r in p.read_text(encoding="utf-8-sig").splitlines()]
    else:
        import openpyxl
        ws = openpyxl.load_workbook(p, data_only=True, read_only=True).worksheets[0]
        rows = [str(r[0]).strip() if r and r[0] is not None else "" for r in
                ws.iter_rows(values_only=True)]
    return [r for r in rows if r]


# Row labels of the "injected activity" sheet -> the Animal field they fill.
_INJ_ROWS = [
    (("id",), "id"), (("alias",), "aliases"), (("weight",), "weight_g"),
    (("act", "full"), "full_mbq"), (("time", "full"), "full_time"),
    (("act", "empty"), "empty_mbq"), (("time", "empty"), "empty_time"),
    (("time", "injection"), "inj_time"),
]


def read_animals(path) -> list[Animal]:
    """Import the lab's 'injected activity' layout: labels down column B, animals across.

    Column A carries the block name (Syringe full / Syringe empty / Injection) and column B
    the field, so 'act (MBq)' is disambiguated by whichever block it sits under.  Rows that
    match nothing land in `extra`, which is what makes the ARRIVE fields survive a round trip.
    """
    import openpyxl
    ws = openpyxl.load_workbook(Path(path), data_only=True, read_only=True).worksheets[0]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    if not rows:
        return []
    ncol = max(len(r) for r in rows)
    ids_row = next((r for r in rows if len(r) > 1 and str(r[1] or "").strip().lower() == "id"), None)
    if ids_row is None:
        raise ValueError("no 'ID' row found in column B — not an injected-activity sheet")
    ids = [str(v).strip() for v in ids_row[2:ncol]]
    animals = [Animal(id=i) for i in ids if i and i.lower() != "none"]
    block = ""
    for r in rows:
        if len(r) < 2:
            continue
        if r[0]:
            block = str(r[0]).strip().lower()
        label = str(r[1] or "").strip().lower()
        if not label:
            continue
        field_name = next((f for keys, f in _INJ_ROWS
                           if all(k in label or k in block for k in keys)), None)
        for i, a in enumerate(animals):
            v = r[2 + i] if len(r) > 2 + i else None
            if v is None or str(v).strip() == "":
                continue
            if field_name == "id":
                continue
            if field_name in ("weight_g", "full_mbq", "empty_mbq"):
                try:
                    setattr(a, field_name, float(str(v).replace(",", ".")))
                except ValueError:
                    pass
            elif field_name in ("full_time", "empty_time", "inj_time"):
                setattr(a, field_name, v.strftime("%H:%M") if hasattr(v, "strftime")
                        else str(v).strip())
            elif field_name == "aliases":
                a.aliases.append(str(v).strip())
            else:
                key = f"{block} {label}".strip() if block and block not in label else label
                a.extra[key] = v.strftime("%H:%M") if hasattr(v, "strftime") else str(v).strip()
    return animals


# --------------------------------------------------------------------- self-check
def _self_check(data_dir=r"C:\Code\BioDist\data_260903"):
    """Rebuild animal 107 of the 260903 experiment from the raw files and check the numbers."""
    d = Path(data_dir)
    tissues = read_column(d / "260903_tissue_list_indiv.xlsx")
    assert tissues[:4] == ["ctrl", "Adrenal", "BAT", "Blood"], tissues[:4]
    assert len(tissues) == 25, len(tissues)

    animals = read_animals(d / "260903_injected_activity.xlsx")
    assert [a.id for a in animals] == ["107", "108", "109", "117", "118", "119",
                                       "110", "111", "120", "121"], [a.id for a in animals]
    a107 = animals[0]
    assert (a107.aliases, a107.weight_g) == (["C1"], 19.6), (a107.aliases, a107.weight_g)
    assert (a107.full_mbq, a107.full_time) == (55.3, "12:12"), (a107.full_mbq, a107.full_time)
    assert (a107.empty_mbq, a107.empty_time) == (6.53, "12:21")
    assert a107.inj_time == "12:17", a107.inj_time
    assert animals[3].extra.get("euth. time") == "15:14", animals[3].extra

    # the rules the reference numbers were checked under (thesis, lab sheets)
    s = Study(name="260903", date="2026-09-03", window="112-168", dt_max=1.5,
              pick_count="auto")
    s.animals = [a107]
    s.tissues = [Tissue(t, "tail" if t == "Tail" else "blank" if t == "ctrl" else "tissue")
                 for t in tissues]
    empty = Source(uid="s1", path=str(d / "Tare-001-20260902-161122-AutoExport.xlsx"),
                   kind="empty", animals=["107"], tissues=tissues)
    filled = Source(uid="s2", path=str(d / "Tare-001-20260903-153146-AutoExport.xlsx"),
                    kind="filled", animals=["107"], tissues=tissues)
    c1 = Source(uid="s3", path=str(d / "Tc-99m-002-20260903-185456-AutoExport.xlsx"),
                kind="count", animals=["107"], tissues=tissues)
    c2 = Source(uid="s4", path=str(d / "Tc-99m-002-20260904-040117-AutoExport.xlsx"),
                kind="count", animals=["107"], tissues=tissues)
    s.sources = [empty, filled, c1, c2]
    # the two paraffin-weighed tissues, and the four vials read on the dose calibrator
    s.manual = [
        Manual("107", "Bone marrow", mass_g=0.00113, note="weighed on paraffin"),
        Manual("107", "Thyr", mass_g=0.00261, note="weighed on paraffin"),
        Manual("107", "Kidneys", mbq=10.85, time="16:46"),
        Manual("107", "caecum", mbq=0.048, time="16:44"),
        Manual("107", "int small", mbq=0.255, time="16:45"),
        Manual("107", "int large", mbq=0.033, time="16:45"),
    ]

    res = compute(s)
    assert res.ref == _dt.datetime(2026, 9, 3, 16, 0), res.ref
    assert [n[:22] for n in res.notes if "control tubes" not in n] == ["107: the tail vial rea"], res.notes

    def pidg(t):
        return res.value(s, "107", t, "pid_g")

    # injected dose: 55.3 MBq @12:12 minus 6.53 @12:21, both decayed to 16:00, tail removed
    assert abs(res.injected_bq["107"] / 1e6 - 31.384) < 0.01, res.injected_bq["107"] / 1e6
    # the tail vial went to the dose calibrator (0.059 MBq, lab sheet): the counter reads background
    assert res.tail_bq["107"] < 100 and any(n.startswith("107: the tail vial reads") for n in res.notes)

    # 107's SPECT started 13:23 (lab sheet): 31.38 MBq at 16:00 taken back 2 h 37 min
    s.animal("107").events = [{"kind": "imaging", "modality": "SPECT/CT", "start": "13:23"},
                              {"kind": "imaging", "modality": "MRI", "start": "14:00"}]
    (mod, t, bq), = at_imaging(s, res, "107")
    assert mod == "SPECT" and abs(bq / 1e6 - 31.384 * 2 ** (157 / 60 / 6.0067)) < 0.05, bq
    s.animal("107").events = []

    # masses: tare difference, and the two hand-entered ones
    assert abs(res.cell("107", "Liver").mass_g - 0.9223) < 1e-4
    assert res.cell("107", "Bone marrow").mass_g == 0.00113
    assert res.cell("107", "Bone marrow").mass_src == "manual"

    # activity: the recount at the lower dead time wins for the liver
    liver = res.cell("107", "Liver")
    assert len(liver.alts) == 2 and liver.dead_time == 1.217, liver
    hi, lo = max(x[1] for x in liver.alts), min(x[1] for x in liver.alts)
    assert abs((hi - lo) / lo * 100 - 4.58) < 0.05, (hi, lo)   # the hot first counting: low
    assert liver.spread_pct(s) is None and not liver.flags, "DT 1.61 is not a valid counting"
    # the default target range (dead time ≤ 1.1) holds neither: the rule among the valid
    # ones, not flagged; valid only up to 1.1 too: none valid — the nearest, flagged
    s.dt_max, s.pick_count = DT_WARN, "all"
    near = compute(s).cell("107", "Liver")
    assert near.dead_time == 1.217 and not near.flags, near.flags
    s.valid_dt = DT_WARN
    near = compute(s).cell("107", "Liver")
    assert near.dead_time == 1.217 and near.flags == ["dead time 1.22"], near.flags
    s.dt_max, s.valid_dt, s.pick_count = 1.5, DT_VALID, "auto"
    c2.kind = "ignored"                                  # kept in the study, left out
    assert len(compute(s).cell("107", "Liver").alts) == 1
    c2.kind = "count"

    # the dose calibrator supplies what the counter could not
    assert res.cell("107", "Kidneys").bq_src == "dose calibrator"
    assert abs(pidg("Kidneys") - 143.2) < 0.5, pidg("Kidneys")
    assert abs(pidg("Blood") - 0.28) < 0.01, pidg("Blood")
    assert abs(pidg("Thyr") - 2.55) < 0.02, pidg("Thyr")
    assert abs(recovery_pct(s, res, "107") - 39.9) < 0.5, recovery_pct(s, res, "107")

    # SUV = (Bq/g in tissue) / (Bq/g whole body); with 19.6 g and 143 %IA/g the kidneys sit ~28
    assert abs(res.value(s, "107", "Kidneys", "suv") - pidg("Kidneys") * 19.6 / 100) < 1e-6

    # %IA/g must not care which instant everything was corrected to
    s.ref_time = "2026-09-03 20:00"
    alt = compute(s)
    assert abs(alt.value(s, "107", "Liver", "pid_g") - pidg("Liver")) < 1e-9, "reference leaked in"
    s.ref_time = ""
    # Bq shown at each animal's injection (12:17 for 107), or at the one study instant
    assert res.ref_of("107") == _dt.datetime(2026, 9, 3, 12, 17)
    assert abs(res.value(s, "107", "Liver", "bq") / res.cell("107", "Liver").bq
               - 2 ** (223 / 60 / 6.00718)) < 1e-6, "16:00 back to 12:17"
    s.ref_rule = "time"
    assert compute(s).value(s, "107", "Liver", "bq") == res.cell("107", "Liver").bq
    s.ref_time = "18:00"                         # a time typed: everything there
    at18 = compute(s)
    assert at18.ref == _dt.datetime(2026, 9, 3, 18) and abs(
        at18.value(s, "107", "Liver", "pid_g") - pidg("Liver")) < 1e-9
    s.ref_rule, s.ref_time = "injection", ""

    # a tail read by hand on the dose calibrator replaces the counted tail vial
    before = res.injected_bq["107"]
    a107.tail_mbq, a107.tail_time = 1.0, "16:00"
    assert abs((before - compute(s).injected_bq["107"]) / 1e6 - (1.0 - res.tail_bq["107"] / 1e6)) < 1e-6
    a107.tail_mbq, a107.tail_time = None, ""

    # a round trip through JSON must land on the same numbers
    again = Study.from_json(s.to_json())
    assert abs(compute(again).value(again, "107", "Liver", "pid_g") - pidg("Liver")) < 1e-9

    # the group run: one file, six vials, tissue-major over three animals
    g = Study(name="group", date="2026-09-03", window="112-168")
    g.animals = [Animal(id=x, isotope="99mTc") for x in ("107", "108", "109")]
    g.tissues = [Tissue("tumour"), Tissue("muscle")]
    gt = Source(uid="g1", path=str(d / "Tare-001-20260902-165551-AutoExport.xlsx"), kind="empty")
    gw = Source(uid="g2", path=str(d / "Tc-99m_weights-010-20260903-141827-AutoExport.xlsx"),
                kind="weigh_count", preset=TISSUE_MAJOR,
                animals=["107", "108", "109"], tissues=["tumour", "muscle"])
    gt.preset, gt.animals, gt.tissues = TISSUE_MAJOR, gw.animals, gw.tissues
    g.sources = [gt, gw]
    gres = compute(g)
    assert abs(gres.cell("107", "tumour").mass_g - 0.2767) < 1e-4
    assert abs(gres.cell("109", "muscle").mass_g - 0.1150) < 1e-4
    assert gres.cell("107", "tumour").bq > gres.cell("107", "muscle").bq

    # animal-major is the same mechanism with the loop the other way round
    gw.preset = gt.preset = ANIMAL_MAJOR
    flipped = compute(g)
    assert abs(flipped.cell("107", "muscle").mass_g - 0.1769) < 1e-4, "107's 2nd vial, not the 4th"

    # the same group run guessed: a second tissue list is a batch counted tissue by tissue,
    # and a six-vial run goes to it, not to the 25-tissue main list
    g.tissues = [Tissue(t) for t in tissues] + [Tissue("tumour", batch="direct"),
                                                 Tissue("muscle", batch="direct")]
    g.sources = [Source(uid="g1", path=gt.path, kind="empty"),
                 Source(uid="g2", path=gw.path, kind="weigh_count")]
    auto_assign(g, load_runs(g))
    assert [x.batch for x in g.sources] == ["direct", "direct"], [x.batch for x in g.sources]
    guessed = compute(g)
    assert abs(guessed.cell("107", "tumour").mass_g - 0.2767) < 1e-4
    assert abs(guessed.cell("109", "muscle").mass_g - 0.1150) < 1e-4

    # the data choice: the first / last counting in range (the first, at dead time 1.61, is
    # valid and in range only up to 2), the widest window of a one-isotope study
    s.pick_count = "first"
    assert compute(s).cell("107", "Liver").bq_src.startswith("Tc-99m-002-20260904")
    s.dt_max = s.valid_dt = 2.0
    assert compute(s).cell("107", "Liver").bq_src.startswith("Tc-99m-002-20260903")
    s.pick_count = "last"
    assert compute(s).cell("107", "Liver").bq_src.startswith("Tc-99m-002-20260904")
    s.pick_count, s.window, s.dt_max, s.valid_dt = "auto", "", 1.5, DT_VALID
    run = load_runs(s)[c1.path]
    assert s.window_for(run).endswith("15-2047"), s.window_for(run)
    s.window_rule = "peak"
    assert s.window_for(run).endswith("112-168"), s.window_for(run)
    s.window, s.window_rule = "112-168", "wide"
    # dates of birth: a cage born over three weeks is a range or a list; its middle ages it
    on = _dt.date(2026, 5, 1)
    assert mid_date("7-26/02/26", on) == mid_date("2026-02-07 – 2026-02-26", on) \
        == _dt.date(2026, 2, 16)
    assert mid_date("07/02, 13/02, 26/02/26", on) == _dt.date(2026, 2, 13)
    assert parse_dates("28/12/25-3/1/26", on)[0][0] == _dt.date(2025, 12, 28)
    assert parse_dates("28/12-3/1/26", on)[0][0] == _dt.date(2025, 12, 28)
    assert format_date("7-26/2/26", on) == "07–26/02/2026", format_date("7-26/2/26", on)
    assert format_date("7/2-3/3/26", on) == "07/02–03/03/2026"
    assert format_date("2026-02-07", on, "DD.MM.YYYY") == "07.02.2026"
    assert format_date("not a date", on) == "not a date"
    assert arrive_gaps(Animal(weight_g=20, extra={"species": "mouse", "sex": "F"}))[:2] == \
        ["strain", "age / DOB"]
    assert [parse_age(x) for x in ("12 wk", "12 weeks", "84 d", "10-12 wk", "12", "x")] == \
        [12, 12, 12, 11, 12, None]
    assert age_weeks("12 wk", _dt.date(2026, 9, 3)) == 12
    assert [colon_time(x) for x in ("12h39", "8 h 05 d1", "2 h p.i.")] ==         ["12:39", "8:05 d1", "2 h p.i."]
    assert parse_time("12h39", _dt.date(2026, 9, 3)) == _dt.datetime(2026, 9, 3, 12, 39)
    # Envigo, 6 wk on arrival 22/9: born 11/8, 7 wk on 30/9
    b = born({"arrival": "22/9/26", "age at arrival": "6 wk"}, _dt.date(2026, 9, 30))
    assert b == _dt.date(2026, 8, 11) and not born({"arrival": "22/9"}, _dt.date(2026, 9, 30))
    assert abs(parse_age("2.5 mo") - 10.87) < 0.01 and abs(parse_age("1 y") - 52.18) < 0.01
    three = ["107", "108", "109"]
    assert strip_summary([(a, t) for t in ("tumor", "muscle") for a in three] + [None]) == \
        "tumor: 107, 108, 109 · muscle: 107, 108, 109 · 1 not used"
    assert strip_summary([("107", t) for t in "ABCDEF"]) == "107: A, B … F"
    assert [names_summary(u, list("ABCDE")) for u in (list("EDCBA"), list("ABD"), ["A"], [])] \
        == ["all", "all but C, E", "A", "none"]
    # tumor 107, 108, 109 then muscle: two vials typed give the order, the rest follows
    six = [None] * 6
    six[0], six[1] = ["107", "tumor"], ["108", "tumor"]
    assert propose(six, 1, three, ["tumor", "muscle"], ANIMAL_MAJOR) == {
        2: ["109", "tumor"], 3: ["107", "muscle"], 4: ["108", "muscle"], 5: ["109", "muscle"]}
    assert propose(six, 0, three, ["tumor", "muscle"], ANIMAL_MAJOR)[1] == ["107", "muscle"]
    assert propose(six, 0, three, ["tumor", "muscle"], TISSUE_MAJOR) == {
        2: ["109", "tumor"], 3: ["107", "muscle"], 4: ["108", "muscle"], 5: ["109", "muscle"]}
    assert propose(six, 1, ["107"], ["tumor"], ANIMAL_MAJOR) == {}, "not on the lists"
    # two empty vials of one name: two tissues, each its own vial and role
    assert unique_names(["a", "(empty)", "(empty)"], {"a"}) == ["a 2", "(empty)", "(empty) 2"]
    assert guess_role("(empty)") == "blank"
    # a rack matched although one tube went: the others agree (260930, animals 5 and 6)
    def rk(*g):
        return (hidex.Run("", "", "", "tare", _dt.datetime(2026, 9, 30, len(g) and 12), None,
                          [], {}, []), [hidex.Slot(2, v + 1, tare_g=x) for v, x in enumerate(g)])
    e1, f1 = rk(5.9, 5.8, 5.7, 5.2), rk(5.95, 5.9, 5.75, 2.5)
    f1[0].started = _dt.datetime(2026, 9, 30, 15)
    assert _match_racks([e1, f1]) == [(0, 1)], "one tube out, the rack still matched"
    f2 = rk(5.95, 5.6, 5.75, 2.5)                # two tubes lighter: not the same rack
    f2[0].started = f1[0].started
    assert _match_racks([e1, f2]) == []
    # an animal's tubes re-ordered by hand in its weighing carry to its counting
    sl = [hidex.Slot(1, v) for v in (1, 2, 3)]
    rn = {p: hidex.Run(p, "", "", "count", None, None, [], {}, sl) for p in ("w", "c", "o")}
    st = Study(sources=[Source("s1", "w", "empty", animals=["2"], tissues=["GB", "K", "L"],
                               slotmap={"1:1": ["2", "K"], "1:2": ["2", "L"],
                                        "1:3": ["2", "GB"]}, auto=False),
                        Source("s2", "c", animals=["2"], tissues=["GB", "K", "L"]),
                        Source("s3", "o", animals=["3"], tissues=["GB", "K", "L"])])
    assert carry_fixes(st, rn) == [("s2", "2", "w", {"1:1": ["2", "K"], "1:2": ["2", "L"],
                                                      "1:3": ["2", "GB"]})], carry_fixes(st, rn)
    st.sources[1].slotmap.update(carry_fixes(st, rn)[0][3])
    assert carry_fixes(st, rn) == [], "once placed the same way, nothing left to carry"
    # grouped by molecule: each group where its first animal is, the order kept inside
    st = Study(animals=[Animal("1", molecule="A"), Animal("2", molecule="B"),
                        Animal("3", molecule="A")], group_by="molecule")
    assert [a.id for a in st.output_animals()] == ["1", "3", "2"]
    st.animal_output = ["2", "3"]
    assert [a.id for a in st.output_animals()] == ["2", "3"]
    # the results / report list: its own order, a rename follows, a gone tissue drops out
    st = Study(tissues=[Tissue("a"), Tissue("b"), Tissue("c", role="blank")])
    assert [t.name for t in st.output_tissues()] == ["a", "b"]
    st.output = ["c", "b", "gone"]
    st.rename_tissue("b", "B")
    assert [t.name for t in st.output_tissues()] == ["c", "B"]
    st.tissue_labels = {"B": "Bone"}
    st.rename_tissue("B", "b2")
    assert st.tissue_label("b2") == "Bone", "a label for the report follows a rename"
    # a procedure brings its kind's own fields, some only for one modality; what a row
    # already holds is never hidden; old name lists come typed
    procs = default_procedures()
    names = [f["name"] for f in ev_fields({"kind": "imaging", "modality": "CT", "kV": "50"},
                                          procs)]
    assert names == ["modality", "start", "duration", "kV", "note"], names
    spect = [f["name"] for f in ev_fields({"kind": "imaging", "modality": "SPECT"}, procs)]
    assert "collimator" in spect and "time per frame" in spect, spect
    assert [f["name"] for f in ev_fields({"kind": "odd", "what": "x"}, procs)] == \
        OTHER_EVENT + ["note"]
    tf = typed_fields()
    assert (tf["euthanasia time"]["type"], tf["supplier"]["type"], "arrival" in tf) == (
        "time", "list", False), tf
    old = typed_procedures({"tumour": ["cells", "date"], "imaging": ["modality", "start",
                                                                     "anaesthesia"]})
    assert [(f["name"], f["type"]) for f in old["tumour"] + old["imaging"]] == [
        ("cells", "text"), ("date", "date"), ("modality", "list"), ("start", "time")], old
    assert event_text({"kind": "imaging", "modality": "SPECT", "start": ""}) \
        == "imaging: modality SPECT"
    assert Source(slotmap={"1:1": ["107", ""]}).mapping(
        hidex.Run(Path(""), "", "", "count", None, None, [], {}, [hidex.Slot(1, 1)])) == {}, \
        "half-typed: no vial placed"

    # an animal's empty tubes, weighed filled, then weighed again: both weighings are its own
    def _w(h, g):
        return hidex.Run(Path(""), "", "", "tare", _dt.datetime(2026, 9, 1, h), None, [], {},
                         [hidex.Slot(r, 1, tare_g=1 + r / 10 + g) for r in (1, 2)])
    e, f, again = _w(1, 0), _w(2, 0.05), _w(3, 0.049)
    blk = [(r, [s]) for r in (e, f, again) for s in r.slots]
    assert sorted(_match_racks(blk)) == [(0, 2), (0, 4), (1, 3), (1, 5)], "reweigh kept"

    # control tubes read 0.02 % light in the filled weighing: every tube of that file is
    def _r(kind, ctrl, tis):
        return hidex.Run(Path(""), "", "", "tare", None, None, [], {},
                         [hidex.Slot(1, 1, tare_g=ctrl), hidex.Slot(1, 2, tare_g=tis)])
    dr = Study(date="2026-09-30", animals=[Animal("1")],
               tissues=[Tissue("ctrl", "blank"), Tissue("Thyroids")],
               sources=[Source(uid=k, path=k, kind=k, animals=["1"], tissues=["ctrl", "Thyroids"],
                               auto=False) for k in ("empty", "filled")])
    druns = {"empty": _r("empty", 6.0, 5.8), "filled": _r("filled", 6.0 * 0.9998, 5.802 * 0.9998)}
    got = compute(dr, druns)
    assert abs(got.cell("1", "Thyroids").mass_g - (5.802 * 0.9998 - 5.8)) < 1e-9
    assert any("control tubes weigh -1.2 mg (-0.020 %)" in n for n in got.notes), got.notes
    dr.drift_fix = "scale"
    got = compute(dr, druns)
    assert abs(got.cell("1", "Thyroids").mass_g - 0.002) < 1e-9, "scaled back by the controls"
    dr.drift_fix = "offset"
    assert abs(compute(dr, druns).cell("1", "Thyroids").mass_g - 0.002) < 5e-5, "offset back"
    dr.drift_fix = "scale"
    range_flags(dr, got, [["Thyroid", 0.5, 4, 0, 0]])
    assert not got.cell("1", "Thyroids").flags
    assert same_tissue("Thyr", "Thyroids") and not same_tissue("Liv", "Liver")
    range_flags(dr, got, [["Thyroid", 0.5, 1.5, 0, 0]])
    assert got.cell("1", "Thyroids").flags == ["out of range: 2 mg (0.5–1.5)"]

    # a sheet of masses typed by hand: an animal's block of empty / full / difference
    art = d.parent / "data_artificial" / "art_weight-TheraSen_260903_biod_24E8.xlsx"
    if art.exists():
        rows, unsure, _ = read_weights(art, [Animal("107", aliases=["C1"])], tissues)
        m = {t: g for a, t, g in rows if a == "107"}
        assert len(m) == 25 and m["Liver"] == 0.9223 and m["Thyr"] == 0.00261 and not unsure, m
    assert read_weights(d / "260903_tissue_list_indiv.xlsx", [Animal("107")], tissues) is None

    # a study keeps its files' data: it computes the same from them alone; its json has a
    # fingerprint that tells an edit made outside
    kept = Study.from_json(s.to_json())
    kept.embedded = {x.path: hidex.pack(hidex.read(x.path)) for x in s.sources}
    for x in kept.sources:
        x.path = "gone/" + Path(x.path).name
    kept.embedded = {"gone/" + Path(k).name: v for k, v in kept.embedded.items()}
    again = Study.from_json(kept.to_json())
    assert not again.outside_edit
    assert abs(compute(again).value(again, "107", "Liver", "pid_g") - pidg("Liver")) < 1e-9
    assert Study.from_json(kept.to_json().replace('"name": "260903"', '"name": "x"')).outside_edit
    assert '"tissues": ["ctrl", "Adrenal"' in kept.to_json(), "short lists on one line"
    assert stamp(c1.path)["sha256"] == hashlib.sha256(Path(c1.path).read_bytes()).hexdigest()

    # studies saved with a single alias still open; new IDs follow the last one
    assert Animal.from_dict({"id": "107", "alias": "C1"}).label == "107 (C1)"
    assert [next_id("S1", set()), next_id("107", {"108"}), next_id("M09", set()),
            next_id("", {"1"})] == ["S2", "109", "M10", "2"]
    assert round(age_weeks("01.06.2026", _dt.date(2026, 9, 3))) == 13
    assert age_weeks("2026-06-01", _dt.date(2026, 9, 3)) == age_weeks("01/06/2026", _dt.date(2026, 9, 3))
    assert age_weeks("1/6/26", _dt.date(2026, 9, 3)) == age_weeks("2026-06-01", _dt.date(2026, 9, 3))
    # date of birth, age, arrival + age at arrival: any one gives the others
    on = _dt.date(2026, 9, 3)
    lf = life_dates({"arrival": "13/8/26", "age at arrival": "6 wk"}, on)
    assert lf["DOB"] == _dt.date(2026, 7, 2) and lf["age"] == 9 and lf["age at arrival"] == 6
    lf = life_dates({"age": "10 wk", "arrival": "6/8/26"}, on)
    assert lf["DOB"] == _dt.date(2026, 6, 25) and lf["age at arrival"] == 6, lf
    assert life_dates({"DOB": "12 wk"}, on)["age"] == 12, "an age typed in the DOB field"
    # a procedure's day: a date, the age then, or D-n
    dob = _dt.date(2026, 6, 25)
    for typed in ("20/8/26", "8 wk", "D-14"):
        assert when_views(typed, dob, on) == {"date": _dt.date(2026, 8, 20), "age": 8,
                                              "day": -14}, typed
    # another loss of the dose (a cotton on the tail) comes off the injected activity
    lo = Study(date="2026-09-03", animals=[Animal.from_dict(
        {"id": "1", "full_mbq": 10, "full_time": "10:00", "loss_mbq": 0.5, "loss_time": "10:00"})])
    lo.animals[0].losses.append({"mbq": "0,5", "time": "10:00", "what": "cotton"})
    got = compute(lo, {})
    assert abs(got.injected_bq["1"] - decay(9e6, _dt.datetime(2026, 9, 3, 10), got.ref,
                                            HALF_LIFE_S["99mTc"])) < 1, "two losses off"
    # an imaging session's anaesthesia, a field until 2026.10.5.1, is a procedure of its own
    old = Animal.from_dict({"id": "1", "events": [
        {"kind": "imaging", "modality": "SPECT/CT", "start": "13:10", "anaesthesia": "iso"}]})
    assert [(e["kind"], e.get("modality"), e.get("agent")) for e in old.events] == [
        ("imaging", "SPECT", None), ("imaging", "CT", None), ("anaesthesia", None, "iso")]
    # a dual-modality session is two procedures
    assert [e["modality"] for e in split_modality([{"kind": "imaging", "modality": "SPECT/CT",
                                                    "start": "13:10"}])] == ["SPECT", "CT"]
    assert "anaesthesia (a procedure, for the imaging or surgery)" in arrive_gaps(
        Animal(events=[{"kind": "imaging"}]))

    # times: bare on the study day, or with a date in any usual order
    day = _dt.date(2026, 9, 23)
    want = _dt.datetime(2026, 9, 24, 8, 46, 37)
    assert parse_time("24/9/26 8:46:37", day) == want
    assert parse_time("8:46:37 9/24", day) == parse_time("2026-09-24 08:46:37", day) == want
    assert parse_time("24.09 8:47", day) == _dt.datetime(2026, 9, 24, 8, 47)
    assert parse_time("14:18", day) == _dt.datetime(2026, 9, 23, 14, 18)
    assert parse_time("24 Sep 8:47", day) is parse_time("25:00", day) is None
    assert parse_time("8:46:37 d1", day) == parse_time("D+1 08:46:37", day) == want
    assert parse_time("d 08:46", day) is None
    assert parse_time("8:46:37.5", day).microsecond == 500000
    assert [format_time(x, day) for x in ("8:46:37 d1", "14:18:59.9", "22/9 23:10", "?")]         == ["08:46 d+1", "14:18", "23:10 d-1", "?"]
    assert format_time("d+1 8:46:37", day, "DD/MM HH:MM") == "24/09 08:46"
    assert set(time_problems(Animal(full_time="14:11", inj_time="14:18", empty_time="14:05",
                                    tail_time="12:13"), day)) == {"empty_time", "tail_time"}
    assert not time_problems(Animal(full_time="14:11", inj_time="14:18", empty_time="14:44",
                                    tail_time="24/9 9:00"), day)

    assert half_life_s("Tc-99m") == half_life_s("99mTc") == half_life_s("\u2079\u2079\u1d50Tc")

    # old studies: kind=tare + role still load; a rename follows the tissue and its role
    old = Source.from_dict({"uid": "s1", "kind": "tare", "role": "filled", "pair": "s0"})
    assert (old.kind, old.auto) == ("filled", False), old
    r = Study(tissues=[Tissue("Control tube", "blank")], manual=[Manual("1", "Control tube")],
              sources=[Source(tissues=["Control tube"], slotmap={"1:1": ["1", "Control tube"]})])
    r.rename_tissue("Control tube", "Gall bladder")
    assert (r.tissues[0].role, r.manual[0].tissue, r.sources[0].tissues,
            r.sources[0].slotmap["1:1"][1]) == ("tissue", "Gall bladder", ["Gall bladder"],
                                                "Gall bladder")

    # guessing: 107's second count is a recount of the first
    a = Study(date="2026-09-03", animals=[Animal("107"), Animal("108")],
              tissues=[Tissue(t) for t in tissues],
              sources=[Source(uid=x.uid, path=x.path, kind=x.kind) for x in (c1, c2)])
    why = {}
    auto_assign(a, load_runs(a), why)
    assert a.sources[0].animals == a.sources[1].animals == ["107"], [x.animals for x in a.sources]
    assert why[a.sources[1].uid].startswith("a recount of") and "107" in why[a.sources[0].uid], why

    # 260923: six animals, 12 weighings, 6 counts; the post-collection weighing put V2's and
    # V3's 4-tube racks into the run before their own
    d = d.parent / "data"
    if (d / "Tubes.txt").exists():
        v = Study(date="2026-09-23", animals=[Animal(x) for x in
                                              ("V1", "V2", "V3", "C1", "C2", "C3")],
                  tissues=[Tissue(t) for t in read_column(d / "Tubes.txt")],
                  mass_rule="first", pick_count="auto", dt_max=1.5)   # as checked then
        v.sources = [Source(uid=f"s{i}", path=str(p)) for i, p in
                     enumerate(sorted(d.glob("*AutoExport.xlsx")))]
        vruns = load_runs(v)
        for x in v.sources:
            x.kind = kind_of(vruns[x.path])
        vnotes = auto_assign(v, vruns)
        assert [x.kind for x in v.sources[:12]] == ["empty"] * 6 + ["filled"] * 6
        # the counts, and the weigh+count recounts of the same tubes the next afternoon
        assert [x.animals for x in v.sources[12:]] == [[x] for x in ("V1", "V2", "V3", "C1",
                                                                     "C2", "C3")] * (
            (len(v.sources) - 12) // 6)
        m = {x.path[-22:-16]: x.mapping(vruns[x.path]) for x in v.sources}
        assert m["170752"]["3:1"] == ("V2", v.tissues[10].name), m["170752"]
        assert m["171857"]["2:4"] == ("V3", v.tissues[13].name)
        assert m["172735"]["1:10"] == ("V3", v.tissues[9].name)
        assert sum("rack slip" in n for n in vnotes) == 2, vnotes
        vres = compute(v, vruns)
        assert all(c.mass_g is not None and c.mass_g > -0.005 for c in vres.cells.values()
                   if c.mass_src), "every tube found its own empty weight"
        # a counting picked by hand for one cell wins over the rule
        k = next(k for k, c in vres.cells.items() if len(c.alts) > 1 and k[1] != "Kidney")
        last = max(vres.cells[k].alts, key=lambda x: x[4])[0]   # the rule: the most counts
        v.chosen = [[*k, "count", last]]
        assert compute(v, vruns).cells[k].bq_src == last != vres.cells[k].bq_src
        alts = vres.cells[k].alts                     # two ticked: weighted, or their mean
        v.chosen = [[*k, "count", *(x[0] for x in alts)]]
        got = compute(v, vruns).cells[k]
        want = sum(x[1] * x[3] for x in alts) / sum(x[3] for x in alts)
        assert got.bq_src == f"weighted of {len(alts)}" and abs(got.bq - want) < 1e-6 * want
        v.combine = "mean"
        got = compute(v, vruns).cells[k]
        want = statistics.fmean(x[1] for x in alts)
        assert got.bq_src == f"mean of {len(alts)}" and abs(got.bq - want) < 1e-6 * want, got.bq
        v.combine = "weighted"
        # a counting in the other window, picked for one cell
        w2 = next(x for x in vres.cells[k].every if x[0] == alts[0][0]
                  and x[1] != alts[0][1])
        v.chosen = [[*k, "count", f"{w2[0]}@{w2[5]}"]]
        got = compute(v, vruns).cells[k]
        assert abs(got.bq - w2[1]) < 1e-6 * w2[1] and "·" in compute(v, vruns).source_label(
            *k, "activity"), got.bq_used
        v.chosen = []
        # two rounds of six files: the day's counting, the next afternoon's recount; a round
        # picked for every cell, or for some
        assert [len(r) for r in vres.rounds] == [6, 6], vres.rounds
        v.pick_count = "round:" + vres.rounds[0][3]
        assert all(vres.round_of(c.bq_src) == 0 for c in compute(v, vruns).cells.values()
                   if c.alts and c.bq_src != "dose calibrator")
        v.pick_count, v.chosen = "auto", [[*k, "count", vres.rounds[1][0]]]
        assert vres.round_of(compute(v, vruns).cells[k].bq_src) == 1
        v.chosen = []
        # the Kidney row closed in the grid: its typed values kept, the counter used again
        kid = next(t for t in v.tissues if t.name == "Kidney")
        kid.closed = ["activity"]
        assert compute(v, vruns).cell("V1", "Kidney").bq_src != "dose calibrator"
        kid.closed = []
        # the recount the next day left the kidneys out (bare vials): the day's weighing stands
        assert all(c.mass_src.startswith("Tare-") for c in vres.cells.values() if c.mass_src)
        bare = [n for n in vres.notes if "no tube in that vial" in n]   # one line for all six
        assert len(bare) == 1 and bare[0].startswith("Kidney of V1 ("), vres.notes
    # efficiencies: the file's own, or the same typed (CPM / 60 / efficiency) — one answer;
    # none at all: CPM only, and said
    s.file_eff = False
    s.efficiency = {k: ["99mTc", w, e] for k, _, w, e in s.eff_rows(load_runs(s))}
    typed = compute(s)
    liver = (res.cell("107", "Liver").bq, typed.cell("107", "Liver").bq)
    assert abs(liver[1] / liver[0] - 1) < 0.002, liver
    s.efficiency = {}
    bare = compute(s)
    assert bare.cell("107", "Liver").bq is None and bare.cell("107", "Liver").raw, "CPM kept"
    assert bare.no_eff == ["Hidex AMG 2240360|⁹⁹ᵐTc_112-168"], bare.no_eff
    s.file_eff = True

    # weighings differ pairwise (13.png, 260930 lymph nodes of 4: 10.5, 8.4, 8.2 mg): the
    # day-of one is out of the consensus, and told against both; the two that agree are not
    # told of it. Two that disagree, no consensus: each told of the other
    ln = Cell(empties=[("tare", 5.9552)], mass_empty="tare",
              fulls=[("w", 5.9657, "filled"), ("r2", 5.9636, "total"), ("r3", 5.9634, "total")])
    got = pairs_of(ln, Study())
    assert [len(got.get(("mass", n), [])) for n in ("w", "r2", "r3")] == [3, 0, 0], got
    assert got[("mass", "w")][0].startswith("out of the consensus — 2 of 3"), got
    ln.fulls = ln.fulls[:2]
    got = pairs_of(ln, Study())
    assert [len(got.get(("mass", n), [])) for n in ("w", "r2")] == [1, 1], got
    # 15.png, 260930 6/Gall bladder: round 3 counted 219 in 15-2047, 57 of them the tissue's
    # (the background off): ±26 %. 32 % under round 1 is 1.8 σ — they agree; and a counting
    # under the valid counts tells nothing to the others
    thin = hidex.Slot(1, 3, secs=30, dead_time=1.0, counts={"w": 219}, cpm={"w": 115})
    assert abs(rsd(thin, "w") - math.sqrt(219) / 57.5) < 1e-9
    assert counts_agree(0.836, 0.03, 0.570, rsd(thin, "w"), Study())
    gb = Cell(every=[("r1", 0.836, 1.0, 3545, None, "w", 3545, True, 1944, 0.03),
                     ("r3", 0.570, 1.0, 115, None, "w", 115, True, 219, 0.03)])
    assert not pairs_of(gb, Study(min_basis="cpm")), "r3 under 1,000 CPM: not valid"
    assert pairs_of(gb, Study(min_basis="cpm", valid_counts=100)), "valid, ±3 % each: differ"
    tubes = [("w", 0.0105, 2), ("r2", 0.0084, 1), ("r3", 0.0082, 1)]     # … and agree: r2, r3
    same = lambda a, b: agree(a[1], b[1], Study())                     # noqa: E731
    assert [u[0] for u in agreeing(tubes, same)] == ["r2", "r3"]
    assert agreeing(tubes[:2], same) == [], "two that disagree: no majority"
    # the rule stays one weighing, one counting: the first (day-of) — unless it is the one at
    # odds with the others, then the first of those that agree
    cts = [("r1", 100.0, 1.0, 40000, None, 0.005), ("r2", 101.0, 1.0, 20000, None, 0.007),
           ("r3", 80.0, 1.0, 20000, None, 0.007)]
    assert [x[0] for x in agreeing(cts, lambda a, b: counts_agree(a[1], a[5], b[1], b[5],
                                                                  Study()))] == \
        ["r1", "r2"], "a counting 20 % off, at ±0.7 %: left out"

    # a time typed after the injection, or as a clock time: each gives the other
    inj = _dt.datetime(2026, 9, 30, 11, 38)
    assert [parse_pi(x) for x in ("2 h p.i.", "1h30 p.i.", "45 min", "45", "13:10")] == [
        _dt.timedelta(hours=2), _dt.timedelta(minutes=90), _dt.timedelta(minutes=45), None, None]
    assert time_views("2 h p.i.", inj, inj.date())["time"] == _dt.datetime(2026, 9, 30, 13, 38)
    assert format_pi(time_views("13:10", inj, inj.date())["pi"]) == "1 h 32 min p.i."
    old = Animal.from_dict({"id": "1", "extra": {"tumour": "TS/A-pc", "anaesthesia": "iso"}})
    assert old.extra == {} and [e["kind"] for e in old.events] == ["tumour", "anaesthesia"]
    old = Animal.from_dict({"id": "1", "extra": {"diet": "chow"}})
    assert old.events == [{"kind": "diet", "what": "chow"}], old.events
    old = Animal.from_dict({"id": "1", "extra": {"anaesthesia": "iso 2 %"},
                            "events": [{"kind": "anaesthesia", "agent": "isoflurane"}]})
    assert old.events == [{"kind": "anaesthesia", "agent": "isoflurane",
                           "note": "anaesthesia: iso 2 %"}], old.events
    print("study self-check ok — animal 107 rebuilt from raw files, group run mapped both ways")


if __name__ == "__main__":
    import sys

    _self_check(*(sys.argv[1:2] or []))
